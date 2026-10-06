# -*- coding: utf-8 -*-
"""SKTC 중국법인 설립 Raw File Data 워크북 생성 (openpyxl).

사용: python build_cn_xlsx.py out.xlsx
계산 칸은 모두 Excel 수식. 미확인 값은 빈칸(=견적 필요)으로 두고 0으로 처리하지 않는다.
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

import cn_data as D

OUT = sys.argv[1]
wb = Workbook()
wb.remove(wb.active)

F_HDR = PatternFill("solid", fgColor="404040")
F_IN = PatternFill("solid", fgColor="DDEBF7")
F_ASK = PatternFill("solid", fgColor="FFF2CC")
F_KEY = PatternFill("solid", fgColor="FDE9E7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {D.S_OFF: "E2EFD9", D.S_EXP: "EAF1FB", D.S_MED: "F2F2F2", D.S_INT: "FFF2CC", D.S_ASM: "FCE4D6",
               D.S_ASK: "FDE9E7", D.S_NA: "D9D9D9"}
TABLES = []


def sheet(name, title, note, tab):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=13, color="EB002C")
    ws["A2"] = note; ws["A2"].font = Font(size=9, color="595959")
    ws["A3"] = f"조사 기준일 {D.BASE_DATE:%Y-%m-%d} · 단위·상태는 각 열 머리글 참조"; ws["A3"].font = Font(size=9, color="808080")
    return ws


def table(ws, name, headers, rows, widths, start=5, wrap_cols=(), num_fmt=None, input_cols=()):
    r0 = start
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.fill = F_HDR; c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            c.border = BOX
            c.alignment = Alignment(wrap_text=(j in wrap_cols), vertical="top")
            if num_fmt and j in num_fmt:
                c.number_format = num_fmt[j]
            if isinstance(v, str) and v in STATUS_FILL:
                c.fill = PatternFill("solid", fgColor=STATUS_FILL[v])
            if j in input_cols and v not in (None, ""):
                c.font = Font(color="0000FF", bold=True); c.fill = F_IN
            if isinstance(v, str) and v.startswith("http"):
                c.hyperlink = v; c.font = Font(color="0563C1", underline="single")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + max(len(rows), 1)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t); TABLES.append(name)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2)
    return r0


def name(n, ws, cell):
    wb.defined_names[n] = DefinedName(n, attr_text=f"'{ws.title}'!${cell[0]}${cell[1:]}")


# ============================================================ 01_README
ws = sheet("01_README", "SKTC 중국법인 설립 검토 — Raw File Data", "원자료 → 법적·사업적 판단 → 계산·비교 → PPT 결론을 ID로 추적하는 워크북", "EB002C")
info = [
    ("목적", "TC 중국법인의 운영 구조·입지(강소성 우선)·인허가·비용·일정 의사결정 근거 보관 및 갱신"),
    ("조사 기준일", f"{D.BASE_DATE:%Y-%m-%d} (내부 회의일). 작업일도 동일"),
    ("수정일 / 버전", f"{D.BASE_DATE:%Y-%m-%d} / v1.0"),
    ("조사 한계", "이 작업 환경의 네트워크 정책으로 중국 정부 사이트 원문을 직접 열지 못함. 웹 검색 결과 요약으로 내용을 확인했으며, 공식 출처라도 상태를 '공식 출처·요약 확인'으로 표시(원문 조문 대조 필요)"),
    ("ID 체계", "SRC(출처) → EVD(근거) → 판단 시트(06·07·08·10) / CITY(도시)·SITE(시설)·MAT(제품)·ASM(가정)·CST(비용)·T(일정)·R(리스크)"),
    ("입력 셀", "파란 글자·연파랑 칸 = 바꿀 수 있는 입력값 (13_FXAssumptions, 08 점수, 16 기간)"),
    ("수식 셀", "검정 글자 = 수식. 계산 결과를 붙여넣은 칸 없음"),
    ("미확인 값", "빈칸 = 견적 필요·정보 없음. 0으로 처리하지 않으며, 합계 옆에 '미완성(견적 n건)'으로 표시"),
    ("외부 연결", "외부 링크·매크로 없음. 출처 URL은 하이퍼링크(클릭 시 브라우저)"),
    ("사용 방법", "① 13 시트 입력값·14 시트 단가를 견적으로 교체 → 15 합계 자동 갱신 ② 08 점수·가중치 변경 → 순위 갱신 ③ 16 기간 변경 → 일정 갱신"),
]
for i, (k, v) in enumerate(info, 5):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True)
    ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
r = 5 + len(info) + 1
ws.cell(row=r, column=1, value="정보 상태 범례").font = Font(bold=True, color="EB002C"); r += 1
for k, v in D.STATUS_LEGEND:
    c = ws.cell(row=r, column=1, value=k); c.fill = PatternFill("solid", fgColor=STATUS_FILL[k])
    ws.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True); r += 1
r += 1
ws.cell(row=r, column=1, value="시트 구성").font = Font(bold=True, color="EB002C"); r += 1
SHEETS = [("02_ExecutiveSummary", "PPT 핵심 결론·근거 ID"), ("03_SourceRegister", "출처 40건"), ("04_EvidenceRaw", "근거 36건 (원문 핵심 구절·한국어 해석)"),
          ("05_Materials", "제품 입력자료·규제 시나리오 A~D"), ("06_LegalMatrix", "법적 쟁점별 결론"), ("07_PermitPaths", "허가 절차·기간"),
          ("08_CityComparison", "도시 관문·점수·민감도 (수식)"), ("09_SitesWarehouses", "단지·창고 후보"), ("10_OperatingModels", "운영 구조 A~G"),
          ("11_MPComparison", "MP 활용 vs TC 법인"), ("12_LogisticsStaff", "직선거리(수식)·인력"), ("13_FXAssumptions", "환율·단가 입력"),
          ("14_CostsRaw", "비용 원자료 (수식·견적 필요 구분)"), ("15_CostModel", "구조별 합계·KRW 환산·미완성 표시"), ("16_Timeline", "일정 (선행관계 수식, 3시나리오)"),
          ("17_RisksActions", "리스크·조치"), ("18_PPTTraceability", "PPT 페이지 ↔ 근거 ID"), ("19_Glossary_Q", "용어 해설·기관 질문지")]
for s, d in SHEETS:
    ws.cell(row=r, column=1, value=s); ws.cell(row=r, column=2, value=d); r += 1
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 120

# ============================================================ 03_SourceRegister
ws = sheet("03_SourceRegister", "03_SourceRegister | 출처 목록", "확인일은 모두 조사 기준일. '공식여부'는 발행 기관 기준이며, 원문 열람 여부와 별개", "2E75B6")
hdr = ["출처 ID", "문서명", "기관·발행자", "문서 종류", "공식 여부", "적용 지역", "발령일", "시행일", "유효 상태", "원문 URL", "확인일", "관련 페이지·조문", "주의사항", "정보 상태"]
rows = []
for s in D.SOURCES:
    st = D.S_OFF if s[4] == "공식" else (D.S_EXP if s[3] in ("전문가 해설",) else D.S_MED)
    rows.append(list(s[:10]) + [f"{D.CHECKED:%Y-%m-%d}"] + list(s[10:]) + [st])
table(ws, "tSource", hdr, rows, [10, 52, 26, 14, 9, 11, 11, 11, 18, 50, 11, 36, 34, 22], wrap_cols=(2, 12, 13))

# ============================================================ 04_EvidenceRaw
ws = sheet("04_EvidenceRaw", "04_EvidenceRaw | 근거 원자료", "원문 핵심 구절은 짧게(검색 요약에서 확인된 문구). '(요약)' 표시는 원문 문구가 아닌 요약", "2E75B6")
hdr = ["근거 ID", "출처 ID", "조사 질문", "원문 핵심 구절·원자료", "한국어 해석", "적용 대상", "적용 조건", "확인·추정 구분", "연결 제품·도시·허가", "PPT 페이지"]
table(ws, "tEvidence", hdr, [list(e) for e in D.EVIDENCE], [10, 10, 26, 48, 48, 18, 26, 20, 18, 10], wrap_cols=(3, 4, 5, 7))

# ============================================================ 05_Materials
ws = sheet("05_Materials", "05_Materials | 제품 입력자료·규제 판단", "제품 정보가 제공되지 않아 CAS·위험등급을 임의로 넣지 않음. 시나리오 A~D로 분석", "FF7900")
hdr = ["제품 ID", "제품 코드·명", "화학명", "CAS", "조성·용매", "IECSC 등재", "신규물질 판단", "위험화학품 판단", "운송 분류·UN", "물성·보관 조건",
       "연간 물량", "기존 등록·备案", "자료 소유권", "정보 상태", "추가 필요 자료", "근거 ID"]
table(ws, "tMaterials", hdr, [list(m) + ["EVD-032, EVD-003"] for m in D.MATERIALS], [9, 26, 10, 9, 10, 16, 16, 18, 14, 22, 10, 16, 12, 12, 46, 16],
      wrap_cols=(2, 10, 15))
r = 5 + len(D.MATERIALS) + 3
ws.cell(row=r, column=1, value="규제 시나리오").font = Font(bold=True, color="EB002C"); r += 1
hdr2 = ["시나리오", "구분", "필요 등록·영업·보관·운송", "사업 영향"]
for j, h in enumerate(hdr2, 1):
    c = ws.cell(row=r, column=j, value=h); c.fill = F_HDR; c.font = Font(bold=True, color="FFFFFF")
for i, sc in enumerate(D.SCENARIOS, 1):
    for j, v in enumerate(sc, 1):
        c = ws.cell(row=r + i, column=j, value=v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BOX
ws.cell(row=r + 6, column=1, value="운송상 위험물(UN 분류)과 위험화학품 목록 해당 여부는 별개로 판정 — 둘 다 확인").font = Font(size=9, color="595959")

# ============================================================ 06_LegalMatrix
ws = sheet("06_LegalMatrix", "06_LegalMatrix | 법적 쟁점 매트릭스", "국가 공통 규정과 성·시 집행 기준 구분. 결론은 근거 기반 판단이며 법률 의견이 아님", "FF7900")
hdr = ["법적 쟁점", "국가·성·시", "관련 규정", "조문", "활동·품목", "의무·예외", "필요 자료", "담당 기관", "결론", "조건", "미확인", "근거 ID"]
table(ws, "tLegal", hdr, [list(x) for x in D.LEGAL], [26, 9, 22, 12, 16, 26, 22, 22, 36, 20, 22, 18], wrap_cols=(1, 6, 7, 9, 10, 11))

# ============================================================ 07_PermitPaths
ws = sheet("07_PermitPaths", "07_PermitPaths | 허가·등록 절차", "현실 예상기간(주)은 분석 가정. 법정 처리기간은 확인된 경우만 기재", "FF7900")
hdr = ["운영 구조", "절차·허가명", "신청 주체", "관할", "선행 조건", "제출 자료", "시설", "인력", "법정 처리기간", "현실 예상기간(주)", "비용", "갱신·유지", "근거 ID"]
table(ws, "tPermit", hdr, [list(x) for x in D.PERMITS], [9, 30, 24, 24, 22, 24, 16, 18, 16, 12, 16, 12, 18], wrap_cols=(2, 3, 4, 5, 6), input_cols=(10,))

# ============================================================ 13_FXAssumptions (먼저 — 이름 정의)
ws = sheet("13_FXAssumptions", "13_FXAssumptions | 환율·단가·가정 입력", "파란 칸만 수정. 빈 칸 = 견적 필요(0 아님)", "70AD47")
hdr = ["가정 ID", "항목", "입력값", "단위", "통화", "적용일", "환율·출처", "가정 이유", "적용 구조", "범위", "확인 상태", "변경 영향"]
table(ws, "tAssume", hdr, [list(a) for a in D.ASSUMPTIONS], [9, 34, 12, 12, 8, 11, 26, 26, 10, 12, 18, 16], wrap_cols=(2, 7, 8), input_cols=(3,))
NAMES = ["FX_KRW", "BASE_WAGE", "WAGE_MULT", "ER_RATE", "LOCAL_HC", "OFFICE_SQM", "OFFICE_RENT", "DEP_MONTHS", "INV_DAYS", "MONTHLY_SALES", "CAP_MULT"]
for i, n in enumerate(NAMES):
    rr = 6 + i
    c = ws.cell(row=rr, column=3)
    if c.value is None:
        c.fill = F_ASK
    name(n, ws, f"C{rr}")
ws["C6"].number_format = "0.00"; ws["C7"].number_format = "#,##0"; ws["C9"].number_format = "0%"

# ============================================================ 14_CostsRaw
ws = sheet("14_CostsRaw", "14_CostsRaw | 비용 원자료 (CNY)", "단가가 빈칸이면 금액도 빈칸(견적 필요). 보증금은 회수 가능 자금으로 초기 지출과 분리", "70AD47")
hdr = ["비용 ID", "항목", "유형", "구조", "단가(CNY)", "수량", "주기", "세금 조건", "근거 구분", "회수 가능", "기준일", "출처·가정 ID", "금액(CNY)", "산출 상태"]
FMAP = {"=사무실월임대료*보증금개월": "=IF(ISNUMBER(OFFICE_RENT),OFFICE_RENT*OFFICE_SQM*DEP_MONTHS,\"\")",
        "=현지기본급*급여배수*(1+사업주부담률)": "=BASE_WAGE*WAGE_MULT*(1+ER_RATE)",
        "=사무실월임대료*12": "=IF(ISNUMBER(OFFICE_RENT),OFFICE_RENT*OFFICE_SQM*12,\"\")",
        "=월판매액*재고일수/30": "=IF(ISNUMBER(MONTHLY_SALES),MONTHLY_SALES*INV_DAYS/30,\"\")",
        "=현지인원": "=LOCAL_HC"}
rows = []
for k, c in enumerate(D.COSTS):
    rr = 6 + k
    price = FMAP.get(c[4], c[4]); qty = FMAP.get(c[5], c[5]) if isinstance(c[5], str) else c[5]
    amt = f'=IF(AND(ISNUMBER(E{rr}),ISNUMBER(F{rr})),E{rr}*F{rr},"")'
    stt = f'=IF(M{rr}="","{c[8] if c[8] != "분석 가정" else "견적 필요"}","산출")'
    rows.append([c[0], c[1], c[2], c[3], price, qty, c[6], c[7], c[8], c[9], f"{D.BASE_DATE:%Y-%m-%d}", c[10], amt, stt])
table(ws, "tCost", hdr, rows, [9, 40, 9, 6, 14, 8, 8, 16, 18, 9, 11, 16, 14, 16], wrap_cols=(2,), num_fmt={5: "#,##0", 13: "#,##0"})
NC = len(D.COSTS); CR = f"$6:${5 + NC}"

# ============================================================ 15_CostModel
ws = sheet("15_CostModel", "15_CostModel | 구조별 비용 합계 (B안 최소 구조 기준)", "확인 가능 금액만 합산하고 빈칸 수를 함께 표시 — 부분합을 전체 예산으로 읽지 말 것", "70AD47")
hdr = ["구분", "유형 키", "확인 가능 합계(CNY)", "견적 필요 건수", "산출 상태", "KRW 환산(백만원)", "비고"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=5, column=j, value=h); c.fill = F_HDR; c.font = Font(bold=True, color="FFFFFF"); c.border = BOX
CM = [("초기 일회성 비용", "초기", "보증금·재고 제외"), ("보증금 (회수 가능)", "보증금", "현금 지출이나 회수 가능"), ("연간 고정비", "연간", "인건비·임대·회계 등"),
      ("물량 연동 변동비 (연)", "변동", "3PL·운송 — 물량 미제공"), ("운전자본 (재고)", "운전자본", "판매액 미제공")]
for i, (lab, key, note) in enumerate(CM, 6):
    ws.cell(row=i, column=1, value=lab); ws.cell(row=i, column=2, value=key)
    ws.cell(row=i, column=3, value=f"=IF(COUNTIFS('14_CostsRaw'!$C$6:$C${5 + NC},B{i},'14_CostsRaw'!$D$6:$D${5 + NC},\"B\")=D{i},\"확인분 없음\",SUMIFS('14_CostsRaw'!$M$6:$M${5 + NC},'14_CostsRaw'!$C$6:$C${5 + NC},B{i},'14_CostsRaw'!$D$6:$D${5 + NC},\"B\"))")
    ws.cell(row=i, column=4, value=f"=COUNTIFS('14_CostsRaw'!$C$6:$C${5 + NC},B{i},'14_CostsRaw'!$D$6:$D${5 + NC},\"B\",'14_CostsRaw'!$M$6:$M${5 + NC},\"\")")
    ws.cell(row=i, column=5, value=f'=IF(D{i}>0,"미완성(견적 "&D{i}&"건)","완성")')
    ws.cell(row=i, column=6, value=f'=IF(ISNUMBER(C{i}),C{i}*FX_KRW/1000000,"-")')
    ws.cell(row=i, column=7, value=note)
    for j in range(1, 8):
        ws.cell(row=i, column=j).border = BOX
    ws.cell(row=i, column=3).number_format = "#,##0"; ws.cell(row=i, column=6).number_format = "#,##0.0"
ws["A12"] = "인건비 시나리오 (1인당 연 인건비 × 인원)"; ws["A12"].font = Font(bold=True, color="EB002C")
hdr2 = ["시나리오", "현지 인원", "1인당 연 인건비(CNY)", "연 인건비(CNY)", "KRW(백만원)", "구성", "상태"]
for j, h in enumerate(hdr2, 1):
    c = ws.cell(row=13, column=j, value=h); c.fill = F_HDR; c.font = Font(bold=True, color="FFFFFF")
SCN = [("최소 (B안)", "=LOCAL_HC", "안전관리자·영업·경리 + 법인장(주재원 별도)"), ("기준 (B안+물류)", "=LOCAL_HC+1", "+수입·물류·통관"),
       ("확대 (C안)", "=LOCAL_HC+3", "+물류 1·창고·품질 2")]
for i, (lab, hc, comp) in enumerate(SCN, 14):
    ws.cell(row=i, column=1, value=lab); ws.cell(row=i, column=2, value=hc)
    ws.cell(row=i, column=3, value="=BASE_WAGE*WAGE_MULT*(1+ER_RATE)").number_format = "#,##0"
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}").number_format = "#,##0"
    ws.cell(row=i, column=5, value=f"=D{i}*FX_KRW/1000000").number_format = "#,##0.0"
    ws.cell(row=i, column=6, value=comp); ws.cell(row=i, column=7, value="분석 가정 (평균임금 기반)")
ws["A18"] = "등록자본 제안 (가정)"; ws["A18"].font = Font(bold=True, color="EB002C")
ws["A19"] = "확인 가능 연간 고정비 × 배수"; ws["C19"] = "=C8*CAP_MULT"; ws["C19"].number_format = "#,##0"
ws["E19"] = '=IF(D8>0,"미완성 — 임대·회계·주재원 견적 반영 후 재산정","산정")'
ws["F19"] = "=C19*FX_KRW/1000000"; ws["F19"].number_format = "#,##0.0"
ws["G19"] = "법정 최저자본 아님. 5년 내 납입(공사법). 등록자본·투자비·보증금·재고자금은 서로 다른 개념"
ws["A21"] = "손익분기 산식 (매출·마진 미제공 → 산식만)"; ws["A21"].font = Font(bold=True, color="EB002C")
ws["A22"] = "필요 연 매출총이익 = 연간 고정비 + 변동비 (견적 반영 후)"; ws["A23"] = "필요 연 매출 = 필요 매출총이익 ÷ 매출총이익률 (내부 입력)"
for col, w in zip("ABCDEFG", (30, 10, 20, 14, 26, 14, 50)):
    ws.column_dimensions[col].width = w

# ============================================================ 08_CityComparison
ws = sheet("08_CityComparison", "08_CityComparison | 도시 관문·평가 점수", "필수 관문 미통과·정보 부족 후보는 점수를 매기지 않음(점수 보류). 점수는 분석 가정(1~5)", "EB002C")
ws["A4"] = "가중치 (합 100, 의사결정용 제안 — 공식 기준 아님)"; ws["A4"].font = Font(bold=True)
for j, cn in enumerate(D.CRIT):
    ws.cell(row=4, column=12 + j, value=cn).font = Font(bold=True, size=9)
WROWS = {}
for k, (wn, ws_) in enumerate(D.WEIGHTS.items()):
    rr = 4 + 1 + k  # 5,6,7 — 표 위에 겹치지 않도록 열 L~Q에 둔다
hdr = ["도시 ID", "성·직할시", "지급시", "구·현·현급시", "운영 구조 역할", "위도", "경도", "G1 법인 설립", "G2 무저장 경영 경로", "G3 적격 창고 후보",
       "점수: " + D.CRIT[0], "점수: " + D.CRIT[1], "점수: " + D.CRIT[2], "점수: " + D.CRIT[3], "점수: " + D.CRIT[4], "점수: " + D.CRIT[5],
       "종합(기본)", "종합(속도 우선)", "종합(비용 우선)", "근거 충족", "판정", "이유", "근거 ID"]
start = 12
rows = []
for k, c in enumerate(D.CITIES):
    rr = start + 1 + k
    sc = list(c[10])
    def tot(wrow):
        return (f'=IF(OR(ISNUMBER(SEARCH("미통과",J{rr})),COUNT(K{rr}:P{rr})<6),"점수 보류",'
                f'SUMPRODUCT(K{rr}:P{rr},$L${wrow}:$Q${wrow})/5)')
    rows.append(list(c[:10]) + sc + [tot(7), tot(8), tot(9), c[11], c[12], c[13], c[14]])
# 가중치 표 (L6:Q9)
ws["K6"] = "가중치 세트"; ws["K6"].font = Font(bold=True)
for j, cn in enumerate(D.CRIT):
    ws.cell(row=6, column=12 + j, value=cn).font = Font(bold=True, size=9)
for k, (wn, wv) in enumerate(D.WEIGHTS.items()):
    ws.cell(row=7 + k, column=11, value=wn)
    for j, w in enumerate(wv):
        c = ws.cell(row=7 + k, column=12 + j, value=w); c.font = Font(color="0000FF", bold=True); c.fill = F_IN
    ws.cell(row=7 + k, column=18, value=f"=SUM(L{7 + k}:Q{7 + k})")
ws["A4"] = None
for j in range(12, 18):
    ws.cell(row=4, column=j, value=None)
table(ws, "tCity", hdr, rows, [10, 12, 14, 22, 26, 7, 7, 8, 18, 24, 7, 7, 7, 7, 7, 7, 10, 10, 10, 8, 20, 40, 22], start=start,
      wrap_cols=(4, 5, 9, 10, 22), input_cols=(11, 12, 13, 14, 15, 16))
for k in range(len(D.CITIES)):
    for col in (17, 18, 19):
        ws.cell(row=start + 1 + k, column=col).number_format = "0.0"
ws.freeze_panes = "E13"
CITY_FIRST, CITY_LAST = start + 1, start + len(D.CITIES)

# ============================================================ 02_ExecutiveSummary
ws = wb.create_sheet("02_ExecutiveSummary", 1)
ws.sheet_properties.tabColor = "EB002C"
ws["A1"] = "02_ExecutiveSummary | 핵심 결론 (PPT 02·28장)"; ws["A1"].font = Font(bold=True, size=13, color="EB002C")
ws["A2"] = "숫자 칸은 다른 시트 수식 참조. 결론은 조건부이며 '확정 불가' 항목은 그대로 표시"; ws["A2"].font = Font(size=9, color="595959")
hdr = ["구분", "내용", "조건·전제", "근거 충족 상태", "PPT 페이지", "근거 ID"]
ES = [
    ("권고 운영 구조", "B안: TC 독자 법인 + 무저장 위험화학품 경영허가 + 허가 외부 창고(3PL)", "제품이 위험화학품·신규물질(시나리오 A)일 때 기준", "중 (규정 근거 기반, 기관 확인 전)", "02, 08, 22", "EVD-014, EVD-015"),
    ("강소성 1순위", "무석 신오구 (법인·영업) + 무석 관할 인정 화공원구 창고", "창고 운영사의 품목·물성 수용", "중", "15, 18", "EVD-014, EVD-020"),
    ("강소성 차선", "쑤저우 SIP 법인 + 장자강 扬子江国际化学工业园 창고", "쑤저우 세칙·창고 품목 확인", "중", "16, 18", "EVD-021"),
    ("인근 대안 1·2", "상하이(수입항·보세 물류) / 허페이(CXMT 대응)", "강소성 창고 미확보 시", "하", "20, 21", "SRC-031, EVD-029"),
    ("연말 기한", "备案 접수 중단(2026-08-15)·기존 备案은 2026-12-31까지 등록증 취득", "TC 제품이 备案으로 수입돼 왔을 때", "중 (공식 통지, 원문 대조 필요)", "05, 26", "EVD-003"),
    ("등록 주체", "현행: 해외 생산자+대리인 가능 / 개정 초안: 중국 생산·수입기업만", "개정 최종 공포 여부", "중", "05, 10", "EVD-001, EVD-002"),
    ("생산설비 필요 가설", "단순 수입·판매에는 생산허가·생산설비 불필요 (근거 기반)", "충진·소분 등 가공 없음", "중", "05, 07", "EVD-009, EVD-011"),
]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=4, column=j, value=h); c.fill = F_HDR; c.font = Font(bold=True, color="FFFFFF")
for i, row in enumerate(ES, 5):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BOX
i = 5 + len(ES)
KF = [("B안 초기 일회성 비용 (확인 가능분, CNY)", "='15_CostModel'!C6", "='15_CostModel'!E6"),
      ("B안 연간 고정비 (확인 가능분, CNY)", "='15_CostModel'!C8", "='15_CostModel'!E8"),
      ("B안 연 인건비 최소 시나리오 (CNY)", "='15_CostModel'!D14", "분석 가정"),
      ("무석 종합점수 (기본 가중치)", f"='08_CityComparison'!Q{CITY_FIRST}", "분석 가정 점수"),
      ("기준 시나리오 초도 공급 예상일", "='16_Timeline'!M20", "분석 가정 일정")]
for k, (lab, f, st) in enumerate(KF):
    ws.cell(row=i + k, column=1, value=lab).font = Font(bold=True)
    c = ws.cell(row=i + k, column=2, value=f); c.number_format = "#,##0"
    ws.cell(row=i + k, column=3, value=st)
ws.cell(row=i + 4, column=2).number_format = "yyyy-mm-dd"
ws.cell(row=i + 3, column=2).number_format = "0.0"
for col, w in zip("ABCDEF", (24, 60, 34, 26, 12, 22)):
    ws.column_dimensions[col].width = w

# ============================================================ 09_SitesWarehouses
ws = sheet("09_SitesWarehouses", "09_SitesWarehouses | 단지·창고 후보", "공개 근거로 확인된 것은 '인정 화공원구' 지위까지. 위험품 창고·품목·특수 물성 수용은 모두 문의 예정", "EB002C")
hdr = ["시설 ID", "도시 ID", "단지·창고 정식 명칭", "주소", "운영사", "시설 종류", "공식 인정·허가", "취급 가능 품목", "특수 물성 수용", "외자·무역법인 입주",
       "임차 조건", "공개 연락처", "문의 상태", "견적 필요", "근거 ID"]
table(ws, "tSite", hdr, [list(s) for s in D.SITES], [9, 9, 40, 14, 12, 16, 16, 22, 16, 14, 12, 14, 10, 8, 14], wrap_cols=(3, 8))

# ============================================================ 10_OperatingModels
ws = sheet("10_OperatingModels", "10_OperatingModels | 운영 구조 A~G", "법적 가능성은 근거 기반 판단(기관 확인 전). 비용·기간은 상대 비교", "FF7900")
hdr = ["구조 ID", "구조", "등록 주체", "수입자", "판매자", "창고 운영자", "필요 허가", "법적 가능성", "선행 조건", "IP·자료", "책임", "투자·운영비", "기간",
       "고객 승인", "장점", "단점", "근거 ID"]
table(ws, "tModel", hdr, [list(m) for m in D.MODELS], [7, 26, 22, 14, 12, 14, 30, 24, 20, 16, 18, 10, 12, 14, 22, 22, 18],
      wrap_cols=(2, 3, 7, 8, 9, 15, 16))

# ============================================================ 11_MPComparison
ws = sheet("11_MPComparison", "11_MPComparison | MP 활용 vs TC 독자 법인", "MP 핵심 정보 미제공 → 시나리오 1(MP 허가 보유)·2(MP 신규 취득)로 비교", "FF7900")
hdr = ["비교 항목", "MP 현황", "TC 신규 법인 조건", "시나리오1: MP 허가 보유", "시나리오2: MP 신규 취득", "수수료·책임", "비용·기간 차이", "데이터 상태", "내부 확인 질문"]
table(ws, "tMP", hdr, [list(m) for m in D.MP], [20, 18, 20, 20, 20, 22, 20, 14, 28], wrap_cols=(1, 4, 5, 6, 9))

# ============================================================ 12_LogisticsStaff
ws = sheet("12_LogisticsStaff", "12_LogisticsStaff | 거리(직선)·인력", "직선거리 = 좌표(도시·구 중심, 대략) 기준 수식. 도로거리·운송 리드타임은 운송사 확인 필요", "70AD47")
hdr = ["도시 ID", "도시", "위도", "경도", "직선거리 CXMT(km)", "직선거리 SK하이닉스 우시(km)", "직선거리 YMTC(km)", "도로거리·시간", "위험물 운송 리드타임"]
cust = {c[0]: (c[3], c[4]) for c in D.CUSTOMERS}
rows = []
for k, c in enumerate(D.CITIES):
    rr = 6 + k
    dist = []
    for cid in ("CUST-01", "CUST-02", "CUST-03"):
        la, lo = cust[cid]
        dist.append(f"=ROUND(6371*ACOS(MIN(1,SIN(RADIANS(C{rr}))*SIN(RADIANS({la}))+COS(RADIANS(C{rr}))*COS(RADIANS({la}))*COS(RADIANS({lo}-D{rr})))),0)")
    rows.append([c[0], f"{c[2]} {c[3]}", f"=INDEX('08_CityComparison'!$F${CITY_FIRST}:$F${CITY_LAST},MATCH(A{rr},'08_CityComparison'!$A${CITY_FIRST}:$A${CITY_LAST},0))",
                 f"=INDEX('08_CityComparison'!$G${CITY_FIRST}:$G${CITY_LAST},MATCH(A{rr},'08_CityComparison'!$A${CITY_FIRST}:$A${CITY_LAST},0))"]
                + dist + ["확인 필요", "확인 필요"])
table(ws, "tDist", hdr, rows, [10, 30, 8, 8, 12, 14, 12, 14, 16])
r = 6 + len(rows) + 2
ws.cell(row=r, column=1, value="고객 위치 (공개 주소 기준)").font = Font(bold=True, color="EB002C"); r += 1
for c in D.CUSTOMERS:
    for j, v in enumerate(c, 1):
        ws.cell(row=r, column=j, value=v)
    r += 1
r += 1
ws.cell(row=r, column=1, value="인력 역할 (최소안 vs 확대안)").font = Font(bold=True, color="EB002C"); r += 1
hdr2 = ["역할", "법정·운영 구분", "자격·교육", "겸직·외주", "상주", "현지/주재원", "최소안(명)", "확대안(명)", "근거 ID"]
for j, h in enumerate(hdr2, 1):
    c = ws.cell(row=r, column=j, value=h); c.fill = F_HDR; c.font = Font(bold=True, color="FFFFFF")
s0 = r + 1
for k, s in enumerate(D.STAFF):
    for j, v in enumerate(s, 1):
        ws.cell(row=s0 + k, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
s1 = s0 + len(D.STAFF)
ws.cell(row=s1, column=1, value="합계").font = Font(bold=True)
ws.cell(row=s1, column=7, value=f"=SUM(G{s0}:G{s1 - 1})"); ws.cell(row=s1, column=8, value=f"=SUM(H{s0}:H{s1 - 1})")
ws.cell(row=s1 + 1, column=1, value="회의 구상 4명(이동선 총감·신규 채용·최영진·경리)은 전담 안전관리자 1명이 포함될 때 무저장(B안) 최소 운영 가능 — 법정 최소 인원 아님")

# ============================================================ 16_Timeline
ws = sheet("16_Timeline", "16_Timeline | 일정 (2026-10-06 착수, 주 단위 · 달력일 기준)", "선행 작업의 최대 종료일 다음 날 시작. 공휴일 미반영 → 영업일 일정 아님. 기간은 분석 가정", "7F7F7F")
ws["B4"] = "착수일"; ws["C4"] = D.BASE_DATE; ws["C4"].number_format = "yyyy-mm-dd"; ws["C4"].font = Font(color="0000FF", bold=True); ws["C4"].fill = F_IN
name("T0", ws, "C4")
hdr = ["작업 ID", "단계", "구조", "책임 부서", "산출물", "선행1", "선행2", "기간 낙관(주)", "기간 기준(주)", "기간 보수(주)", "시작(기준)", "종료(기준)", "종료(기준)·참조용",
       "종료(낙관)", "종료(보수)", "기간 근거", "법정·현실", "병목", "시작(낙관)", "시작(보수)"]
rows = []
n = len(D.TASKS)
for k, t in enumerate(D.TASKS):
    rr = 6 + k
    preds = [p.strip() for p in t[5].split(",")] if t[5] != "-" else []
    p1 = preds[0] if preds else ""; p2 = preds[1] if len(preds) > 1 else ""

    def start(endcol):
        rng = f"${endcol}$6:${endcol}${5 + n}"
        a = f'IF(F{rr}="",T0-1,INDEX({rng},MATCH(F{rr},$A$6:$A${5 + n},0)))'
        b = f'IF(G{rr}="",T0-1,INDEX({rng},MATCH(G{rr},$A$6:$A${5 + n},0)))'
        return f"=MAX({a},{b})+1"
    rows.append([t[0], t[1], t[2], t[3], t[4], p1, p2, t[7], t[8], t[9],
                 start("L"), f"=K{rr}+I{rr}*7-1", f"=L{rr}",
                 f"=S{rr}+H{rr}*7-1", f"=T{rr}+J{rr}*7-1", t[10], t[11], t[12], start("N"), start("O")])
table(ws, "tTask", hdr, rows, [7, 40, 7, 18, 16, 7, 7, 9, 9, 9, 11, 11, 11, 11, 11, 24, 10, 16, 11, 11], wrap_cols=(2,), input_cols=(8, 9, 10))
for k in range(n):
    for col in (11, 12, 13, 14, 15, 19, 20):
        ws.cell(row=6 + k, column=col).number_format = "yyyy-mm-dd"
TROW = {t[0]: 6 + k for k, t in enumerate(D.TASKS)}

# 02 시트 참조 셀 보정 (T14 종료(기준))
wb["02_ExecutiveSummary"].cell(row=5 + len(ES) + 4, column=2, value=f"='16_Timeline'!L{TROW['T14']}")

# ============================================================ 17_RisksActions
ws = sheet("17_RisksActions", "17_RisksActions | 리스크·확인 과제", "기한은 착수 후 2~6주 내 제안 일정. 진행 상태는 모두 '문의 예정'", "7F7F7F")
hdr = ["리스크 ID", "내용", "영향", "현재 근거", "확인 대상", "담당 부서", "우선순위", "대응", "기한", "진행 상태", "완료 증거", "PPT 페이지"]
table(ws, "tRisk", hdr, [list(x) + ["문의 예정", "-", "27"] for x in D.RISKS], [8, 44, 8, 22, 20, 12, 8, 40, 11, 10, 10, 9], wrap_cols=(2, 8))

# ============================================================ 18_PPTTraceability
ws = sheet("18_PPTTraceability", "18_PPTTraceability | PPT ↔ 근거 추적", "PPT 주요 주장과 숫자의 근거 시트·ID", "7F7F7F")
TRACE = [
    ("02", "경영진 요약", "B안·무석 조건부 추천 / 연말 기한 / 등록 주체", "B안 비용(확인 가능분)", "02_ExecutiveSummary, 15_CostModel", "EVD-002, EVD-003, EVD-014", "SRC-004, SRC-005, SRC-011", "ASM-001~011", "조건부", "근거 기반·기관 확인 전"),
    ("05", "회의 가설 검증", "5개 주장 검증 결과", "-", "06_LegalMatrix", "EVD-001~011", "SRC-001, SRC-005, SRC-009", "-", "개정안 확정 여부", "요약 확인"),
    ("10", "화학물질 등록·IP", "신청 주체·시험자료·정보보호", "-", "06_LegalMatrix", "EVD-001~007", "SRC-001~008, SRC-036, SRC-040", "-", "-", "요약 확인"),
    ("11", "위험화학품 허가 경로", "무저장·임차저장·저장 / 처리기한", "20 영업일(사례)", "07_PermitPaths", "EVD-008~018", "SRC-009~011, SRC-035", "-", "무석 기한 확인", "요약 확인"),
    ("13~18", "강소성 후보", "원구 23곳·무석·쑤저우·창저우", "점수 76/69/63", "08_CityComparison, 09_SitesWarehouses", "EVD-020, EVD-021", "SRC-013, SRC-030", "점수 가정", "점수=분석 가정", "근거 기반"),
    ("20~22", "인근 대안·민감도", "상하이·허페이 / 가중치 3세트", "점수 68 등", "08_CityComparison", "SRC-031~034", "-", "가중치", "-", "근거 부족"),
    ("23", "MP vs TC", "시나리오 1·2", "-", "11_MPComparison", "EVD-002, EVD-005", "-", "-", "MP 정보 없음", "내부 확인 필요"),
    ("25", "비용", "초기·연간·인건비·등록자본", "15 시트 값", "15_CostModel, 14_CostsRaw, 13_FXAssumptions", "EVD-027, EVD-034, EVD-036", "SRC-024~026", "ASM-001~011", "견적 필요 다수", "미완성"),
    ("26", "일정", "기준 시나리오·연말 대응", "16 시트 값", "16_Timeline", "EVD-003, EVD-016", "SRC-005, SRC-035", "기간 가정", "공휴일 미반영", "분석 가정"),
    ("27", "리스크·2주 과제", "R01~R13", "-", "17_RisksActions", "-", "-", "-", "-", "문의 예정"),
]
table(ws, "tTrace", ["PPT 페이지", "슬라이드", "핵심 주장", "숫자·비교 결과", "원자료 시트", "근거 ID", "출처 ID", "가정 ID", "조건·한계", "검증 상태"], TRACE,
      [9, 18, 36, 18, 30, 24, 26, 12, 16, 16], wrap_cols=(3, 5, 6, 7))

# ============================================================ 19_Glossary_Q
ws = sheet("19_Glossary_Q", "19_Glossary_Q | 용어 해설·질문지", "부록 용어표·기관 질문지 원본", "BFBFBF")
table(ws, "tGloss", ["용어", "쉬운 뜻", "이번 사업과의 관계"], [list(g) for g in D.GLOSSARY], [36, 60, 46], wrap_cols=(1, 2, 3))
r = 5 + len(D.GLOSSARY) + 3
ws.cell(row=r, column=1, value="질문지 (진행 상태: 모두 문의 예정)").font = Font(bold=True, color="EB002C"); r += 1
for grp, qs in D.QUESTIONS.items():
    ws.cell(row=r, column=1, value=grp).font = Font(bold=True); r += 1
    for q_ in qs:
        ws.cell(row=r, column=2, value=q_).alignment = Alignment(wrap_text=True)
        ws.cell(row=r, column=3, value="문의 예정"); r += 1

# 시트 순서 정리
order = ["01_README", "02_ExecutiveSummary", "03_SourceRegister", "04_EvidenceRaw", "05_Materials", "06_LegalMatrix", "07_PermitPaths",
         "08_CityComparison", "09_SitesWarehouses", "10_OperatingModels", "11_MPComparison", "12_LogisticsStaff", "13_FXAssumptions", "14_CostsRaw",
         "15_CostModel", "16_Timeline", "17_RisksActions", "18_PPTTraceability", "19_Glossary_Q"]
wb._sheets = [wb[s] for s in order]
wb.save(OUT)
print("saved", OUT, len(wb.sheetnames), "sheets")
