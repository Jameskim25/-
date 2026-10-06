# -*- coding: utf-8 -*-
"""PPT 편집 도우미 — 기존 SK trichem 양식 도형의 서식을 유지하며 텍스트·위치를 바꾸고, 공정도를 넓힌다."""
import copy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Noto Sans KR"
E = 914400


def shapes(slide):
    return {sh.name: sh for sh in slide.shapes}


def set_par(p, texts):
    """paragraph의 기존 run 서식을 유지하며 텍스트 교체. texts: str 또는 run 텍스트 리스트."""
    if isinstance(texts, str):
        texts = [texts]
    runs = p.runs
    if not runs:
        r = p.add_run(); r.text = texts[0]
        runs = p.runs
    for i, t in enumerate(texts):
        if i < len(runs):
            runs[i].text = t
        else:
            last = p.runs[-1]._r
            nr = copy.deepcopy(last); last.addnext(nr)
            p.runs[-1].text = t
    for r in p.runs[len(texts):]:
        r._r.getparent().remove(r._r)


def set_lines(sh, lines):
    tf = sh.text_frame
    paras = tf.paragraphs
    while len(paras) < len(lines):
        last = paras[-1]._p
        last.addnext(copy.deepcopy(last))
        paras = tf.paragraphs
    for i, ln in enumerate(lines):
        set_par(paras[i], ln)
    for p in tf.paragraphs[len(lines):]:
        p._p.getparent().remove(p._p)


def pos(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = Emu(int(x * E))
    if y is not None: sh.top = Emu(int(y * E))
    if w is not None: sh.width = Emu(int(max(w, 0.005) * E))
    if h is not None: sh.height = Emu(int(h * E))


def delete(sh):
    sh._element.getparent().remove(sh._element)


def clone(slide, sh):
    el = copy.deepcopy(sh._element)
    slide.shapes._spTree.append(el)
    new = slide.shapes[-1]
    ids = [s.shape_id for s in slide.shapes]
    el.xpath("./*[1]/p:cNvPr")[0].set("id", str(max(ids) + 1))
    return new


def copy_fill(src, dst):
    """src 도형의 채우기·선·글자색을 dst에 복사 (spPr fill/ln, 첫 run 색)."""
    sp_s = src._element.spPr; sp_d = dst._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:ln"):
        for el in sp_d.findall(qn(tag)):
            sp_d.remove(el)
    geom = sp_d.find(qn("a:prstGeom"))
    anchor = geom if geom is not None else sp_d[-1]
    for tag in ("a:ln", "a:solidFill", "a:noFill"):
        el = sp_s.find(qn(tag))
        if el is not None:
            anchor.addnext(copy.deepcopy(el))
    if src.has_text_frame and dst.has_text_frame and src.text_frame.paragraphs[0].runs:
        sr = src.text_frame.paragraphs[0].runs[0]
        for p in dst.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = sr.font.size; r.font.bold = sr.font.bold
                if sr.font.color and sr.font.color.type == 1:
                    r.font.color.rgb = sr.font.color.rgb


def fill(sh, hexcol):
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(hexcol)


def runfmt(r, size, bold=False, color="3F3F3F"):
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color); r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def box(slide, x, y, w, h, color, text="", tsize=5.5, tcolor="FFFFFF", bold=True, line=None, dash=False, align="c"):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(max(w, 0.01)), Inches(h))
    sh.shadow.inherit = False
    if color:
        fill(sh, color)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    tf = sh.text_frame; tf.margin_left = tf.margin_right = Inches(0.01); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = {"c": PP_ALIGN.CENTER, "l": PP_ALIGN.LEFT}[align]
    if text:
        r = p.add_run(); r.text = text; runfmt(r, tsize, bold, tcolor)
    return sh


def tbox(slide, x, y, w, h, runs, align="l", wrap=True):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if isinstance(runs[0], list) else [runs]
    for i, pr in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"c": PP_ALIGN.CENTER, "l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT}[align]
        for t, size, bold, color in pr:
            r = p.add_run(); r.text = t; runfmt(r, size, bold, color)
    return sh


def cell_set(cell, text, tmpl=None):
    tf = cell.text_frame
    p = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if not p.runs:
        r = p.add_run()
        if tmpl is not None and tmpl.text_frame.paragraphs[0].runs:
            r._r.insert(0, copy.deepcopy(tmpl.text_frame.paragraphs[0].runs[0]._r.find(qn("a:rPr"))))
    set_par(p, text)


def table_rows(tbl, rows, tmpl_col=1):
    """rows: list of lists (len = ncols). None = 유지."""
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is None:
                continue
            c = tbl.cell(i, j)
            tmpl = tbl.cell(i, tmpl_col) if j != tmpl_col else tbl.cell(i, 0)
            cell_set(c, v, tmpl)


def unmerge_all(tbl):
    for i in range(len(tbl.rows)):
        for j in range(len(tbl.columns)):
            c = tbl.cell(i, j)
            if c.is_merge_origin:
                c.split()


def normalize_table(tbl, fills=None, mar=27000):
    """분할·병합으로 서식이 빠진 셀에 같은 행 기준 셀(1열)의 tcPr을 복사하고 여백·정렬·문단 끝 크기를 통일."""
    fills = fills or {}
    for i in range(len(tbl.rows)):
        ref = tbl.cell(i, 1)._tc.get_or_add_tcPr()
        for j in range(len(tbl.columns)):
            tc = tbl.cell(i, j)._tc
            tcPr = tc.get_or_add_tcPr()
            if tcPr.find(qn("a:lnB")) is None and j != 1:
                new = copy.deepcopy(ref)
                tc.replace(tcPr, new); tcPr = new
            tcPr.set("anchor", "ctr"); tcPr.set("marL", str(mar)); tcPr.set("marR", str(mar)); tcPr.set("marT", "0"); tcPr.set("marB", "0")
            if (i, j) in fills:
                for f in tcPr.findall(qn("a:solidFill")):
                    tcPr.remove(f)
                sf = etree.SubElement(tcPr, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", fills[(i, j)])
            for para in tbl.cell(i, j).text_frame.paragraphs:
                if j > 0:
                    para.alignment = PP_ALIGN.CENTER
                sz = para.runs[0].font.size if para.runs else None
                end = para._p.find(qn("a:endParaRPr"))
                if end is None:
                    end = etree.SubElement(para._p, qn("a:endParaRPr"))
                end.set("sz", str(int(sz.pt * 100)) if sz else "600")


ICON_TXT = {"정제기", "정제기1", "정제기2", "시차"}


def widen_flow(slide, x_from=0.31, x_old=6.98, x_new=9.69, y0=1.0, y1=3.75):
    """우측 운영조건 패널을 지우고 공정 플로우를 가로로 확대 (아이콘은 비율 유지, 연결선은 아이콘 가장자리에 맞춤)."""
    for sh in list(slide.shapes):
        t, l = sh.top / E, sh.left / E
        if y0 <= t < y1 and l >= 7.0:
            delete(sh)
    f = (x_new - x_from) / (x_old - x_from)
    lin = lambda x: x_from + (x - x_from) * f
    region = [sh for sh in slide.shapes if y0 <= sh.top / E < y1]

    def is_icon(sh):
        n = sh.name
        txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
        return n.startswith("Can") or n.startswith("Round Same Side") or n == "Rectangle 177" or txt in ICON_TXT
    icons = []
    for sh in region:
        if is_icon(sh) and sh.shape_type != 9:
            l, w = sh.left / E, sh.width / E
            icons.append((l, l + w, lin(l + w / 2) - (l + w / 2)))

    def mapx(x):
        for l, r, dx in icons:
            if l - 0.03 <= x <= r + 0.03:
                return x + dx
        return lin(x)
    for sh in region:
        l, w = sh.left / E, sh.width / E
        if sh.shape_type == 9 or sh.name.startswith("Connector"):
            x1, x2 = mapx(l), mapx(l + w)
            pos(sh, x=min(x1, x2), w=abs(x2 - x1))
        elif is_icon(sh):
            pos(sh, x=mapx(l + w / 2) - w / 2)
        else:
            pos(sh, x=lin(l), w=w * f)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

