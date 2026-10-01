# -*- coding: utf-8 -*-
"""PPT 붙여넣기용 가로형 시트 (P1~P5) — 원자료 입력 칸 겸 표.

· 파란 글자 칸 = 입력 칸(원자료). 01_Inputs·02_Batch_Raw·04~06 계산 시트가 이 칸을 참조하므로,
  어느 P 표에서 값을 고쳐도 다른 P 표·계산 시트가 모두 같이 바뀐다.
· 검정 글자 칸 = 수식(다른 표·계산 시트에서 자동 계산).
· 열 폭·글꼴·색을 PPT 표 크기에 맞춰, 범위를 복사해 PPT에 붙여넣으면 같은 크기의 표가 된다.

build_xlsx.py는 이 모듈의 PIN(입력 칸 주소)·ship_cell()·P3 달력 주소를 먼저 읽어 01_Inputs 등에서 참조하고,
마지막에 add_ppt_sheets()로 시트를 만든다.
"""
import datetime as dt

from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.properties import PageSetupProperties

PF = "Noto Sans KR"
DARK, GRAY, INK = "1A1A1A", "7F7F7F", "0000CC"     # INK = 입력 칸 글자색
HX, CXC, EZ, HS = "EB002C", "2E75B6", "FF7900", "7030A0"
line = Side(style="thin", color="D9D9D9")
BOT = Border(bottom=line)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)
LEF = Alignment(horizontal="left", vertical="center", wrap_text=False, shrink_to_fit=True)
KG1 = '#,##0.0;[Color10]-#,##0.0;"0"'
KG0 = '#,##0;-#,##0;"0"'
SGN = '[Color10]+#,##0;[Red]-#,##0;"0"'
H1 = '0.0'

S1, S2, S3, S4, S5 = "P1_공정시간", "P2_일정_Gantt", "P3_과거Batch_2026", "P4_2027_월별", "P5_2028_월별"


def q(sheet):
    return f"'{sheet}'!"


# ------------------------------------------------------------------ 입력 칸 주소 (단일 원천)
PIN = {}            # key -> (sheet, cell, default, number_format)


def _pin(key, sheet, cell, default, nf=None):
    PIN[key] = (sheet, cell, default, nf)


# P1: 공정시간 표 (5 Gal 행이 공통 공정시간 입력) · 충진·용기 조건 · 생산능력
for k, col, v in (("t_iqc", "C", 2), ("t_prep", "D", 2), ("t_ref", "E", 45), ("t_pqctr", "F", 6), ("t_fqc", "G", 2)):
    _pin(k, S1, f"{col}6", v, H1)
_pin("t_oqc5", S1, "J6", 2, H1)
_pin("t_ez_m", S1, "I7", 8, H1)
_pin("t_oqc2", S1, "J7", 2, H1)
_pin("t_hs_m", S1, "I8", None, H1)
for k, col, v, nf in (("t_5g", "C", 2, H1), ("n_5g_b", "D", 9, "0"), ("kg_5g", "E", 20, "0"), ("kg_ez", "F", 140, "0"),
                      ("kg_hs", "G", 150, "0"), ("t_ez_a", "H", None, H1), ("t_hs_a", "I", None, H1), ("kg_b", "J", 190, "0"),
                      ("wh_day", "K", 8, "0")):
    _pin(k, S1, f"{col}13", v, nf)
_pin("max11", S1, "C17", 11, "0")
_pin("capa_now", S1, "I17", 26.2, '0.0" 톤/년"')
# P2: 일정·Gantt 입력 줄 (20칸씩 9개) — 대정비 월 생산량은 P4 4월 생산 계획 칸이 원천
P2_KEYS = [("maint_m", "대정비 월 (매년)", 4), ("maint_kg", "대정비 월 생산 kg (P4 4월)", None),
           ("rf_s", "정제기 2대 생산 시작", dt.date(2027, 5, 1)), ("f12_e", "최초 12개월 종료", dt.date(2028, 4, 30)),
           ("ars_ez", "ARS 시작 (이지켐)", dt.date(2027, 7, 1)), ("ars_hs", "ARS 시작 (한솔)", dt.date(2027, 7, 1)),
           ("t105", "105℃ 50.2톤 적용 시작", dt.date(2028, 1, 1)), ("offset", "정제기 2 기동 시차 h", 30), ("int_rf", "설비별 투입 간격 h", None)]
P2_W = 20
for i, (k, lab, v) in enumerate(P2_KEYS):
    if k == "maint_kg":
        continue
    _pin(k, S2, f"{CL(3 + P2_W * i)}6", v, "yyyy-mm-dd" if isinstance(v, dt.date) else ('0"월"' if k == "maint_m" else "0.0"))
# P4: 2027 생산 계획 (1~4월 확정 · 4월 = 대정비 월 생산량 → Capa. 기준 차감값)
PLAN_ROW = 6
PLAN27_FIX = [2178, 2178, 3168, 1188]
_pin("maint_kg", S4, f"F{PLAN_ROW}", 1188, "#,##0")
# P3: 과거 Batch 달력 (5~12월 × 1~31일) · 월 Total · 2026 출하
P3_MONTHS = list(range(5, 13))
P3_GRID = f"{q(S3)}$C$6:$AG${5 + len(P3_MONTHS)}"
P3_GRID_FIRST = f"{q(S3)}$C$6"
P3_TOTAL = {m: f"{q(S3)}$AI${6 + i}" for i, m in enumerate(P3_MONTHS)}
P3_TOTAL_DEF = {6: 8, 7: 11, 8: 8, 9: 8, 10: 11}
P3_C0 = 38                                  # 2026 표 시작 열 (AL)
P3_SHIP_ROW = {"hx": 13, "ez": 14, "cx": 15, "hs": 16}
P45_SHIP_ROW = {yr: {"hx": 9, "cx": 10, "ez": 11, "hs": 12} for yr in (2027, 2028)}
SHIP_DEF = {
    2026: {"hx": [1600, 1420, 1440, 1580, 1200, 1200, 1360, 1420, 1380, 1480, 1480, 1420], "ez": [0] * 8 + [560] * 4,
           "cx": [None] * 12, "hs": [None] * 12},
    2027: {"hx": ["=22720/12"] * 12, "cx": [0, 0, 580, 580, 580, 580] + [780] * 6, "ez": [560] * 6 + [1120] * 6, "hs": [300] * 12},
    2028: {"hx": ["=22720/12"] * 12, "cx": [780] * 12, "ez": [840] * 12, "hs": [300] * 12}}


def ship_cell(yr, cust, m):
    """연·고객·월 → 입력 칸 절대 참조."""
    if yr == 2026:
        return f"{q(S3)}${CL(P3_C0 + m)}${P3_SHIP_ROW[cust]}"
    return f"{q(S4 if yr == 2027 else S5)}${CL(2 + m)}${P45_SHIP_ROW[yr][cust]}"


def plan_cell(yr, m):
    """연·월 → P4/P5 생산 계획 칸 절대 참조 (2027~2028)."""
    return f"{q(S4 if yr == 2027 else S5)}${CL(2 + m)}${PLAN_ROW}"


def pin_ref(key):
    sheet, cell, _, _ = PIN[key]
    col = "".join(ch for ch in cell if ch.isalpha()); row = "".join(ch for ch in cell if ch.isdigit())
    return f"{q(sheet)}${col}${row}"


# ------------------------------------------------------------------ 서식 도우미
def inch_w(inches):
    px = inches * 96
    return round(px / 12, 3) if px < 12 else round((px - 5) / 7, 2)


def pfont(size=7, bold=False, color=DARK):
    return Font(name=PF, size=size, bold=bold, color=color)


def fill(c):
    return PatternFill("solid", fgColor=c)


def setup(ws, title, note, tab="404040"):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws["A1"] = title; ws["A1"].font = pfont(11, True, "EB002C")
    ws["A2"] = note; ws["A2"].font = pfont(7, False, GRAY)
    ws.column_dimensions["A"].width = 1.5


def section(ws, row, col, text, sub="", sub_col=None):
    ws.cell(row=row, column=col, value=text).font = pfont(8.5, True)
    if sub:
        ws.cell(row=row, column=sub_col or col + 4, value=sub).font = pfont(6.5, False, GRAY)
    ws.row_dimensions[row].height = 15


def put(ws, r, c, v, nf=None, bold=False, color=DARK, bg=None, align="c", size=7):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = pfont(size, bold, color); cell.border = BOT
    cell.alignment = CEN if align == "c" else LEF
    if nf:
        cell.number_format = nf
    if bg:
        cell.fill = fill(bg)
    return cell


def inp(ws, r, c, v, nf=None, size=7, bold=True):
    """입력 칸: 파란 글자."""
    return put(ws, r, c, v, nf, bold=bold, color=INK, size=size)


def header_row(ws, r, c0, labels, bg="404040"):
    for j, lab in enumerate(labels):
        put(ws, r, c0 + j, lab, bold=True, color="FFFFFF", bg=bg, align="l" if j == 0 else "c")
    ws.row_dimensions[r].height = 12


def note(ws, r, c, text):
    cell = put(ws, r, c, text, align="l", color=GRAY, size=6); cell.border = Border()


def _row(C, key):
    return int(C[key].split("$")[-1])


LEGEND = ("파란 글자 = 입력 칸 (여기서 고치면 다른 P 표·계산 시트가 모두 같이 바뀜) · 검정 = 자동 계산. "
          "PPT: 범위 복사 → 붙여넣기(원본 서식 유지/그림) · 계속 연동하려면 [선택하여 붙여넣기 → 연결하여 붙여넣기].")


# ------------------------------------------------------------------ 시트 생성
def add_ppt_sheets(wb, R, C, MAPX, batches):
    M0, SR = MAPX["M0"], MAPX["SR"]
    Mq = "'04_Monthly_2026_2028'!"
    rows_04 = lambda yr, m: M0 + (yr - 2026) * 12 + (m - 1)
    Cq = "'03_Capacity_Model'!"

    # ============================================================ P1
    ws = wb.create_sheet(S1)
    setup(ws, "P1 | As-is 공정시간 · 현재 생산능력 (PPT 1장)", LEGEND, "EB002C")
    for j, w in enumerate([1.7] + [0.85] * 9):
        ws.column_dimensions[CL(2 + j)].width = inch_w(w)
    section(ws, 4, 2, "① 한 Batch 확인 공정시간 (h)", "단순 합계 — 달력 납기·다음 Batch 투입 간격 아님 · 5 Gal 행 = 공통 공정시간 입력")
    header_row(ws, 5, 2, ["경로", "IQC", "준비·투입", "정제", "PQC+이송", "FQC", "충진 전 소계", "충진", "OQC·출하", "확인시간 합계"])
    for r, lab in ((6, "5 Gal 약 9병"), (7, "이지켐 200 L 1용기 (수동)"), (8, "한솔 200 L 1용기")):
        put(ws, r, 2, lab, bold=True, align="l")
        for col in "CDEFG":
            if r == 6:
                inp(ws, r, " CDEFG".index(col) + 2, PIN[{"C": "t_iqc", "D": "t_prep", "E": "t_ref", "F": "t_pqctr", "G": "t_fqc"}[col]][2], H1)
            else:
                put(ws, r, " CDEFG".index(col) + 2, f"={col}$6", H1)
        put(ws, r, 8, f"=SUM(C{r}:G{r})", H1, bold=True, bg="F2F2F2")
    put(ws, 6, 9, "=C13*D13", H1)
    inp(ws, 6, 10, 2, H1); inp(ws, 7, 9, 8, H1); inp(ws, 7, 10, 2, H1); inp(ws, 8, 9, None, H1); put(ws, 8, 10, "=J7", H1)
    put(ws, 6, 11, "=H6+I6+J6", H1, bold=True, color=HX)
    put(ws, 7, 11, "=H7+I7+J7", H1, bold=True, color=HX)
    put(ws, 8, 11, '=IF(ISNUMBER(I8),H8+I8+J8,TEXT(H8+J8,"0")&" + 미확인")', H1, bold=True, color=HX)
    note(ws, 9, 2, "※ 57 h = IQC + 준비·투입 + 정제 + PQC·이송 + FQC (중복 가산 없음) · 5 Gal 충진 = 병당 h × 병수 (아래 입력) · 한솔 수동 충진시간 미확인")
    section(ws, 11, 2, "충진 · 용기 · Batch 조건 (입력)")
    header_row(ws, 12, 2, ["조건", "5 Gal h/병", "병/Batch", "kg/병", "이지켐 kg/용기", "한솔 kg/용기", "이지켐 ARS h", "한솔 ARS h", "kg/Batch", "근무 h/일"])
    put(ws, 13, 2, "값", bold=True, align="l")
    for k in ("t_5g", "n_5g_b", "kg_5g", "kg_ez", "kg_hs", "t_ez_a", "t_hs_a", "kg_b", "wh_day"):
        _, cell, v, nf = PIN[k]
        inp(ws, 13, " ABCDEFGHIJK".index(cell[0]), v, nf)
    section(ws, 15, 2, "② 정제기 1대 생산능력 비교", "월 720 h 기준")
    header_row(ws, 16, 2, ["구분", "과거 최대", "", "정제 + 기타", "", "이론 (정제만)", "", "현재 Capa.", "", ""])
    for a, b in ((3, 4), (5, 6), (7, 8), (9, 11)):
        ws.merge_cells(start_row=16, start_column=a, end_row=16, end_column=b)
    put(ws, 17, 2, "Batch / 월", bold=True, align="l")
    inp(ws, 17, 3, 11, "0")
    put(ws, 17, 5, f'=TEXT(C17*E6,"0")&" h + "&TEXT(720-C17*E6,"0")&" h"', bold=True)
    put(ws, 17, 7, "=720/E6", "0.0", bold=True)
    inp(ws, 17, 9, 26.2, '0.0" 톤/년"')
    put(ws, 18, 2, "kg / 월 · 산식", bold=True, align="l")
    put(ws, 18, 3, "=C17*J13", KG0, bold=True)
    put(ws, 18, 5, "합계 720 h", bold=True)
    put(ws, 18, 7, "=INT(G17)*J13", KG0, bold=True)
    put(ws, 18, 9, '=TEXT(I17*1000/J13,"0.0")&" Batch 상당 · 월 "&TEXT(I17*1000/J13/12,"0.00")', bold=True)
    for r in (17, 18):
        for a, b in ((3, 4), (5, 6), (7, 8), (9, 11)):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    ws.print_area = "B4:K18"

    # ============================================================ P2
    ws = wb.create_sheet(S2)
    setup(ws, "P2 | 운영 일정 · 정제기 2대 시간차 운전 30일 (PPT 2장)", LEGEND, "EB002C")
    NC = 180
    ws.column_dimensions["B"].width = inch_w(1.3)
    for j in range(NC):
        ws.column_dimensions[CL(3 + j)].width = inch_w(7.9 / NC)
    last = CL(2 + NC)
    section(ws, 4, 2, "일정 · 운전 조건 (입력)", "여기 날짜·시간을 바꾸면 아래 일정표·Gantt와 P4·P5·계산 시트가 같이 바뀜", sub_col=3 + 60)
    put(ws, 5, 2, "항목", bold=True, color="FFFFFF", bg="404040", align="l")
    put(ws, 6, 2, "값", bold=True, align="l")
    for i, (k, lab, v) in enumerate(P2_KEYS):
        c = 3 + P2_W * i
        put(ws, 5, c, lab, bold=True, color="FFFFFF", bg="404040", size=6)
        if k == "maint_kg":
            put(ws, 6, c, f"={pin_ref('maint_kg')}", "#,##0", bold=True, size=6.5)
        else:
            inp(ws, 6, c, v, PIN[k][3], size=6.5)
        for rr in (5, 6):
            ws.merge_cells(start_row=rr, start_column=c, end_row=rr, end_column=c + P2_W - 1)
    note(ws, 7, 2, "※ 대정비 월 생산량 = P4 4월 생산 계획 칸 (1,188 kg) · 개선 후 월 Capa. = (연간 Capa. − 대정비 월 생산량) ÷ 11 · "
                   "2028 105℃ 50.2톤 · 설비별 투입 간격 공란 = 월 Capa. 역산 참고값 사용 · 시차는 운영 개념 설명용 가정")
    section(ws, 9, 2, "① 투자 · 운영 일정 (2027.1~2028.6)", "생산 계획 행 = P4·P5 생산 계획 연동 (kg/월)", sub_col=3 + 60)
    put(ws, 10, 2, "구분", bold=True, color="FFFFFF", bg="404040", align="l")
    for j in range(18):
        y, m = (2027, j + 1) if j < 12 else (2028, j - 11)
        c = 3 + j * 10
        put(ws, 10, c, f"=DATE({y},{m},1)", 'yy"."m', bold=True, color="FFFFFF", bg="404040" if y == 2027 else "7F7F7F", size=6)
        ws.merge_cells(start_row=10, start_column=c, end_row=10, end_column=c + 9)
    mm = f'MONTH({{d}})={R("maint_m")}'
    sched = [("■ 현재 정제기 1대 (확정 계획)", f'IF(AND({{d}}<{R("rf_s")},NOT({mm})),"1대 확정","")', "F2F2F2"),
             ("■ 대정비 (매년)", f'IF({mm},"대정비","")', "D9D9D9"),
             ("■ 정제기 2대 개선 생산", f'IF(AND({{d}}>={R("rf_s")},NOT({mm})),IF({{d}}<={R("f12_e")},"12개월","생산"),"")', "FDE9E7"),
             ("■ ARS 200 L 충진", f'IF({{d}}>=MIN({R("ars_hs")},{R("ars_ez")}),"ARS","수동")', "E4DFEC"),
             ("■ 105℃ 50.2톤 기준", f'IF(AND(ISNUMBER({R("t105")}),{{d}}>={R("t105")},NOT({mm})),"105℃","")', "E4DFEC"),
             ("생산 계획 (kg/월)", None, None)]
    r = 11
    for lab, fx, bg in sched:
        put(ws, r, 2, lab, bold=True, align="l")
        for j in range(18):
            c = 3 + j * 10
            y, m = (2027, j + 1) if j < 12 else (2028, j - 11)
            if fx is None:
                put(ws, r, c, f"={plan_cell(y, m)}", "#,##0", bold=True, color=CXC, size=6)
            else:
                put(ws, r, c, "=" + fx.replace("{d}", f"${CL(c)}$10"), size=6)
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 9)
        if fx is not None:
            ws.conditional_formatting.add(f"C{r}:{last}{r}", FormulaRule(formula=[f'AND(C{r}<>"",C{r}<>"수동",C{r}<>"미정")'],
                                                                       fill=fill(bg), font=Font(name=PF, bold=True, color="C00000")))
        r += 1
    r += 1
    section(ws, r, 2, "② 정제기 2대 시간차 운전 개념 — 30일 (720 h)", "1칸 = 4 h · 준비 2 + 정제 + PQC·이송 + FQC 시간은 P1 입력 사용", sub_col=3 + 60)
    r += 1
    SL = 6
    put(ws, r, 2, "일 →", bold=True, color="FFFFFF", bg="404040", align="l")
    for dday in range(30):
        c = 3 + dday * SL
        put(ws, r, c, dday + 1, "0", bold=True, color="FFFFFF", bg="404040", size=6)
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + SL - 1)
    hr2 = r; r += 1
    off, gi = f"{Cq}$B${_row(C, 'off')}", f"{Cq}$B${_row(C, 'g_int')}"
    a0 = R("t_iqc"); a1 = f"({R('t_iqc')}+{R('t_prep')}+{R('t_ref')})"
    i1 = f"({a1}+{R('t_pqctr')}+{R('t_fqc')})"; f1 = f"({i1}+{R('n_5g_b')}*{R('t_5g')})"
    for lab, o, kind in (("정제기 1", "0", "ref"), ("정제기 2", off, "ref"), ("PQC·이송·FQC", None, "ins"), ("5 Gal 충진 (공용)", None, "fil")):
        put(ws, r, 2, lab, bold=True, align="l", size=6.5)
        for j in range(NC):
            t = f"({j}*24/{SL}+2)"
            ph = lambda oo: f"MOD({t}-{oo},{gi})"
            if kind == "ref":
                fx = f"=IF(AND({ph(o)}>={a0},{ph(o)}<{a1}),1,IF({ph(o)}>={a1},2,0))"
            elif kind == "ins":
                fx = f"=IF(OR(AND({ph(0)}>={a1},{ph(0)}<{i1}),AND({ph(off)}>={a1},{ph(off)}<{i1})),3,0)"
            else:
                fx = f"=IF(OR(AND({ph(0)}>={i1},{ph(0)}<{f1}),AND({ph(off)}>={i1},{ph(off)}<{f1})),4,0)"
            ws.cell(row=r, column=3 + j, value=fx).number_format = ";;;"
        ws.row_dimensions[r].height = 11
        r += 1
    for v, col in ((1, "FF7900"), (2, "FFF2CC"), (3, "2E75B6"), (4, "548235")):
        ws.conditional_formatting.add(f"C{hr2 + 1}:{last}{r - 1}", CellIsRule(operator="equal", formula=[str(v)], fill=fill(col)))
    put(ws, r, 2, "누적 생산 kg", bold=True, align="l", size=6.5)
    for dday in range(30):
        c = 3 + dday * SL; T = (dday + 1) * 24
        cnt = "+".join(f"(INT(({T}-{oo}-{i1})/{gi})-INT((0-{oo}-{i1})/{gi}))" for oo in ("0", off))
        put(ws, r, c, f"=({cnt})*{R('kg_b')}", "#,##0", color=CXC, size=5.5)
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + SL - 1)
    r += 1
    note(ws, r, 2, "범례: 주황 = 준비·정제 · 연노랑 = 기타(Mix·이송·대기, 확인) · 파랑 = PQC·이송·FQC · 초록 = 5 Gal 충진")
    ws.print_area = f"B9:{last}{r}"

    # ============================================================ P3
    ws = wb.create_sheet(S3)
    setup(ws, "P3 | 과거 Batch 표시일자 · 간격 · 생산능력 · 2026 출하 (PPT 3장)", LEGEND, "7F7F7F")
    ws.column_dimensions["B"].width = inch_w(0.55)
    for d in range(31):
        ws.column_dimensions[CL(3 + d)].width = inch_w(0.155)
    for j in range(3):
        ws.column_dimensions[CL(34 + j)].width = inch_w(0.42)
    section(ws, 4, 2, "① Batch 표시일자 달력 (입력: 해당 날짜 칸에 Batch 번호)", "→ 02_Batch_Raw 간격·관측 간격·04 과거계획 자동 반영", sub_col=3 + 18)
    put(ws, 5, 2, "월", bold=True, color="FFFFFF", bg="404040", align="l")
    for d in range(31):
        put(ws, 5, 3 + d, d + 1, "0", bold=True, color="FFFFFF", bg="404040", size=5.5)
    for j, lab in enumerate(("번호", "Total", "차이")):
        put(ws, 5, 34 + j, lab, bold=True, color="FFFFFF", bg="404040", size=6)
    pos = {(m, d): n for n, m, d in batches}
    for i, m in enumerate(P3_MONTHS):
        r = 6 + i
        put(ws, r, 2, f"{m}월" + (" *" if m == 5 else ""), bold=True, align="l")
        for d in range(1, 32):
            inp(ws, r, 2 + d, pos.get((m, d)), "0", size=5.5)
        put(ws, r, 34, f"=COUNT(C{r}:AG{r})", "0", bold=True)
        inp(ws, r, 35, P3_TOTAL_DEF.get(m), "0")
        put(ws, r, 36, f'=IF(ISNUMBER(AI{r}),AH{r}-AI{r},"–")', SGN, bold=True)
    g0, g1 = 6, 5 + len(P3_MONTHS)
    ws.conditional_formatting.add(f"C{g0}:AG{g1}", FormulaRule(formula=[f'COLUMN(C{g0})-2>DAY(EOMONTH(DATE({R("hist_y")},ROW(C{g0})-{g0}+5,1),0))'], fill=fill("BFBFBF")))
    ws.conditional_formatting.add(f"C{g0}:AG{g0}", FormulaRule(formula=[f'C{g0}<>""'], fill=fill("7F7F7F"), font=Font(name=PF, bold=True, color="FFFFFF")))
    ws.conditional_formatting.add(f"C{g0 + 1}:AG{g1}", FormulaRule(formula=[f'C{g0 + 1}<>""'], fill=fill(EZ), font=Font(name=PF, bold=True, color="FFFFFF")))
    ws.conditional_formatting.add(f"AH{g0}:AJ{g1}", FormulaRule(formula=[f'AND(ISNUMBER($AJ{g0}),$AJ{g0}<>0)'], font=Font(name=PF, bold=True, color="C00000")))
    note(ws, g1 + 1, 2, "* 5월 일부(#37·#38) · 번호 = 입력된 Batch 수 · Total = 원자료 월 Total(입력) · 빨강 = 불일치 · 회색 칸 = 없는 날짜")
    Bq = "'02_Batch_Raw'!"
    tot = MAPX["SUMROW"]["tot"]
    r = g1 + 3
    section(ws, r, 2, "② 표시일자 간격 분포 (분석 범위)"); r += 1
    put(ws, r, 2, "간격", bold=True, color="FFFFFF", bg="404040", align="l")
    for j, lab in enumerate(("1일", "2일", "3일", "4일", "5일")):
        put(ws, r, 3 + j * 3, lab, bold=True, color="FFFFFF", bg="404040"); ws.merge_cells(start_row=r, start_column=3 + j * 3, end_row=r, end_column=5 + j * 3)
    put(ws, r, 18, "평균 간격", bold=True, color="FFFFFF", bg="404040"); ws.merge_cells(start_row=r, start_column=18, end_row=r, end_column=25)
    put(ws, r, 26, "중앙값", bold=True, color="FFFFFF", bg="404040"); ws.merge_cells(start_row=r, start_column=26, end_row=r, end_column=30)
    r += 1
    put(ws, r, 2, "건수", bold=True, align="l")
    for j, col in enumerate("OPQRS"):
        put(ws, r, 3 + j * 3, f"={Bq}{col}{tot}", "0", bold=True); ws.merge_cells(start_row=r, start_column=3 + j * 3, end_row=r, end_column=5 + j * 3)
    put(ws, r, 18, f"=TEXT({Cq}B{_row(C, 'ob_int')},\"0.0\")&\" h (\"&TEXT({Cq}B{_row(C, 'ob_lo')},\"0.0\")&\"~\"&TEXT({Cq}B{_row(C, 'ob_hi')},\"0.0\")&\")\"", bold=True, color=EZ)
    ws.merge_cells(start_row=r, start_column=18, end_row=r, end_column=25)
    put(ws, r, 26, f"={Cq}B{_row(C, 'ob_med')}&\" h\"", bold=True); ws.merge_cells(start_row=r, start_column=26, end_row=r, end_column=30)
    ws.print_area = f"B4:AJ{r}"
    # 오른쪽 블록: ③ 생산능력 비교 · ⑤ 2026
    c0 = P3_C0
    ws.column_dimensions[CL(c0)].width = inch_w(1.55)
    for j in range(13):
        ws.column_dimensions[CL(c0 + 1 + j)].width = inch_w(0.56)
    section(ws, 4, c0, "③ 월 생산능력 비교 (정제기 1대 · 720 h)")
    header_row(ws, 5, c0, ["기준", "간격 h", "", "월 Batch", "", "월 kg", "", "연 환산 t", "", "26.2 대비", "", "", "", ""])
    for k in range(4):
        rr, cr = 6 + k, MAPX["CMP0"] + k
        put(ws, rr, c0, f"={Cq}A{cr}", bold=True, align="l")
        for j, (col, nf) in enumerate((("B", "0.0"), ("C", "0.0"), ("D", KG0), ("E", "0.00"), ("F", '+0.00;-0.00;"0"'))):
            put(ws, rr, c0 + 1 + j * 2, f"={Cq}{col}{cr}", nf, bold=True, color=HX if k == 2 else DARK)
            ws.merge_cells(start_row=rr, start_column=c0 + 1 + j * 2, end_row=rr, end_column=c0 + 2 + j * 2)
        put(ws, rr, c0 + 11, f"={Cq}H{cr}", align="l", color=GRAY, size=6)
        ws.merge_cells(start_row=rr, start_column=c0 + 11, end_row=rr, end_column=c0 + 13)
    section(ws, 11, c0, "⑤ 2026 출하 제시분 vs 현재 Capa. (kg)", "고객 행 = 입력 · CXMT·한솔 미제시(공란) — 합계 미포함", sub_col=c0 + 5)
    header_row(ws, 12, c0, ["구분"] + [f"{m}월" for m in range(1, 13)] + ["연간"])
    for cust, lab, colr in (("hx", "■ SK하이닉스", HX), ("ez", "■ 이지켐", EZ), ("cx", "■ CXMT", CXC), ("hs", "■ 한솔", HS)):
        rr = P3_SHIP_ROW[cust]
        put(ws, rr, c0, lab, bold=True, align="l", color=colr)
        for m in range(12):
            inp(ws, rr, c0 + 1 + m, SHIP_DEF[2026][cust][m], KG0, bold=False)
        put(ws, rr, c0 + 13, f'=IF(COUNT({CL(c0 + 1)}{rr}:{CL(c0 + 12)}{rr})=0,"미제시",SUM({CL(c0 + 1)}{rr}:{CL(c0 + 12)}{rr}))', KG0, bold=True, bg="F2F2F2")
    for rr, lab, col, nf in ((17, "출하 합계 (제시분)", "Y", KG0), (18, "계획 Total×190", "M", KG0), (19, "계획 번호×190", "L", KG0),
                             (20, "추정 생산 가능 (관측)", "S", KG0), (21, "추정 − 출하", "AP", SGN)):
        put(ws, rr, c0, lab, bold=True, align="l")
        for m in range(1, 13):
            src = f"{Mq}{col}{rows_04(2026, m)}"
            fx = f'=IF(ISNUMBER({src}),{src}*{R("kg_b")},"")' if col in ("L", "M") else f"={src}"
            put(ws, rr, c0 + m, fx, nf, bold=col in ("Y", "S", "AP"))
        a, b = rows_04(2026, 1), rows_04(2026, 12)
        ft = f"=SUM({Mq}{col}{a}:{col}{b})*{R('kg_b')}" if col in ("L", "M") else f"=SUM({Mq}{col}{a}:{col}{b})"
        put(ws, rr, c0 + 13, ft, nf, bold=True, bg="F2F2F2")
    cy = MAPX["CMPY"]["26_now"]
    put(ws, 22, c0, "현재 Capa. 대비", bold=True, align="l")
    put(ws, 22, c0 + 1, f"=TEXT('05_Reflux_Scenarios'!B{cy},\"0.00\")&\"톤 = 현재 \"&TEXT({R('capa_now')},\"0.0\")&\"톤의 \"&TEXT('05_Reflux_Scenarios'!B{cy}/{R('capa_now')},\"0.0%\")&\" · 산술 차이 \"&TEXT('05_Reflux_Scenarios'!D{cy},\"0.00\")&\"톤 (가동률·판매 가능량 아님)\"",
        bold=True, color=HX, align="l")
    ws.merge_cells(start_row=22, start_column=c0 + 1, end_row=22, end_column=c0 + 13)

    # ============================================================ P4 / P5
    S5q = "'05_Reflux_Scenarios'!"
    cy, RES = MAPX["CMPY"], MAPX["RES"]
    for yr, sheet, slide in ((2027, S4, 4), (2028, S5, 5)):
        ws = wb.create_sheet(sheet)
        setup(ws, f"P{slide} | {yr} 월별 생산 계획 · 출하 · 충진 (PPT {slide}장 표) — kg",
              LEGEND + (" 생산 계획 1~4월 + 고객 4행이 입력 칸 (4월 = 대정비 월 생산량 → Capa. 기준 차감값)." if yr == 2027 else " 고객 4행이 출하 입력 칸 · 생산 계획은 P2·P4 입력으로 자동 계산."), "FF7900")
        ws.column_dimensions["B"].width = inch_w(2.1)
        for j in range(12):
            ws.column_dimensions[CL(3 + j)].width = inch_w(0.53)
        ws.column_dimensions["O"].width = inch_w(0.85)
        header_row(ws, 4, 2, ["구분"] + [f"{m}월" for m in range(1, 13)] + ["연간"])
        srr = SR[yr]
        body = [("운전 조건", "D", None, None),
                ("생산 계획 (kg)" + (" — 1~4월 확정" if yr == 2027 else " — 50.2톤 기준"), "PLAN", KG1, None),
                ("Batch 환산 (÷190)", "PB", "0.0", None) if yr == 2027 else ("참고: 47.2톤 기준 (105℃ 미적용)", "P47", KG1, None),
                ("생산 계획 − 출하", "AP", SGN, None), ("■ SK하이닉스", "hx", KG1, HX), ("■ CXMT", "cx", KG0, CXC),
                ("■ 이지켐", "ez", KG0, EZ), ("■ 한솔", "hs", KG0, HS), ("출하 합계", "Y", KG1, None),
                ("필요 Batch (출하÷Batch량)", "Z", "0.0", None), ("5 Gal 충진 h (병 상당×h/병)", "AB", "0.0", None),
                ("이지켐 용기", "AD", "0", None), ("한솔 용기", "AE", "0", None), ("200 L 충진 방식 (이지켐/한솔)", "AFAG", None, None),
                ("200 L 충진 h (확인분+미확인)", "AJX", "0.0", None), ("누적 (생산 계획 − 출하) '27.1~", "AQ", SGN, None)]
        mm = R("maint_m")

        def capa_fx(y, m):
            d0 = f"DATE({y},{m},1)"
            return (f"=IF({m}={mm},{R('maint_kg')},IF(AND(ISNUMBER({R('t105')}),{d0}>={R('t105')}),{R('capa_m105')},"
                    f"IF({d0}>={R('rf_s')},{R('capa_m')},{R('capa_now')}*1000/12)))")
        r = 5
        for lab, col, nf, colr in body:
            put(ws, r, 2, lab, bold=True, align="l", color=colr or DARK)
            if col in ("hx", "cx", "ez", "hs"):
                assert r == P45_SHIP_ROW[yr][col], (yr, col, r)
                for m in range(12):
                    inp(ws, r, 3 + m, SHIP_DEF[yr][col][m], nf, bold=False)
                put(ws, r, 15, f"=SUM(C{r}:N{r})", KG0, bold=True, bg="F2F2F2")
                r += 1
                continue
            for m in range(1, 13):
                rr_ = rows_04(yr, m)
                if col == "PLAN":
                    assert r == PLAN_ROW
                    if yr == 2027 and m <= len(PLAN27_FIX):
                        inp(ws, r, 2 + m, PLAN27_FIX[m - 1], KG0)
                        continue
                    fx = capa_fx(yr, m)
                elif col == "PB":
                    fx = f"={CL(2 + m)}{PLAN_ROW}/{R('kg_b')}"
                elif col == "P47":
                    fx = f"=IF({m}={mm},{R('maint_kg')},{R('capa_m')})"
                elif col == "AFAG":
                    fx = f"={Mq}AF{rr_}&\"/\"&{Mq}AG{rr_}"
                elif col == "AJX":
                    fx = f'=IF(ISNUMBER({Mq}AJ{rr_}),{Mq}AJ{rr_},TEXT(N({Mq}AH{rr_})+N({Mq}AI{rr_}),"0")&"+미확인")'
                else:
                    fx = f"={Mq}{col}{rr_}"
                put(ws, r, 2 + m, fx, nf, bold=col in ("Y", "Z", "AB", "AP", "PLAN", "AQ"), size=6.5 if col in ("D", "AFAG", "AJX") else 7,
                    color=CXC if col == "PLAN" else DARK)
            total = {"D": "-", "AFAG": "-", "AJX": "-", "AQ": f"=N{r}"}.get(col, f"=SUM(C{r}:N{r})")
            put(ws, r, 15, total, {KG1: KG0, "0.0": "#,##0.0"}.get(nf, nf), bold=True, bg="F2F2F2")
            if col == "D":
                ws.conditional_formatting.add(f"C{r}:N{r}", FormulaRule(formula=[f'LEFT(C{r},2)="개선"'], fill=fill("FDE9E7")))
                ws.conditional_formatting.add(f"C{r}:N{r}", FormulaRule(formula=[f'C{r}="대정비"'], fill=fill("D9D9D9"), font=Font(name=PF, bold=True, color="C00000")))
            if col in ("PLAN", "D"):
                ws.conditional_formatting.add(f"C{r}:N{r}", FormulaRule(formula=[f'C$5="대정비"'], fill=fill("EDEDED")))
            if col == "AJX":
                ws.conditional_formatting.add(f"C{r}:O{r}", FormulaRule(formula=[f'NOT(ISNUMBER(C{r}))'], fill=fill("FFF2CC"), font=Font(name=PF, color="C55A11", size=6.5)))
            r += 1
        R5 = lambda k: f"{S5q}$B${RES[k]}"
        if yr == 2027:
            lines = [f"=\"생산 계획 \"&TEXT({R5('plan27')}/1000,\"0.00\")&\"톤 (1~4월 확정 \"&TEXT({R5('plan27_fix')}/1000,\"0.00\")&\" + 5~12월 \"&TEXT({R('capa_m')},\"#,##0.0\")&\" kg × \"&TEXT({R5('n_rf27')},\"0\")&\") vs 출하 \"&TEXT({S5q}B{cy['27_plan']},\"0.00\")&\"톤 → \"&TEXT({S5q}D{cy['27_plan']},\"+0.00;-0.00\")&\"톤 · 누적 최저 \"&TEXT(-{R5('inv_need')}/1000,\"+0.00;-0.00\")&\"톤 (\"&TEXT({R5('inv_low')},\"yy.m\")&\") = 필요 선행재고\"",
                     f"=\"월 Capa. = (\"&TEXT({R('capa_rf')}*1000,\"#,##0\")&\" − 대정비 \"&TEXT({R('maint_kg')},\"#,##0\")&\") ÷ 11 = \"&TEXT({R('capa_m')},\"#,##0.0\")&\" kg · 상반기 출하 \"&TEXT({S5q}B{cy['27_h1']},\"0.00\")&\" vs 계획 \"&TEXT({S5q}C{cy['27_h1']},\"0.00\")&\" · 하반기 \"&TEXT({S5q}B{cy['27_h2']},\"0.00\")&\" vs \"&TEXT({S5q}C{cy['27_h2']},\"0.00\")&\" → \"&TEXT({S5q}D{cy['27_h2']},\"+0.00;-0.00\")&\"톤\""]
        else:
            lines = [f"=\"출하 가정 \"&TEXT({S5q}B{cy['28_plan']},\"0.00\")&\"톤 vs 생산 계획 \"&TEXT({S5q}C{cy['28_plan']},\"0.00\")&\"톤 (105℃ 50.2 · 4월 대정비 + 11개월 \"&TEXT({R('capa_m105')},\"#,##0.0\")&\" kg) → \"&TEXT({S5q}D{cy['28_plan']},\"+0.00;-0.00\")&\"톤 · 47.2 기준이면 \"&TEXT({S5q}D{cy['28_47']},\"+0.00;-0.00\")&\"톤\""]
        for fx in lines:
            put(ws, r, 2, "비교 (t)", bold=True, align="l")
            put(ws, r, 3, fx, bold=True, color=HX, align="l"); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=15)
            r += 1
        note(ws, r, 2, ("* 파란 글자 = 입력 · 생산 계획 5~12월 = Capa. 기준 자동 계산 (대정비 월 = P4 4월 칸)" if yr == 2027 else
                        "* 파란 글자 = 입력 · 생산 계획 = (50.2톤 − 대정비 P4 4월 칸) ÷ 11 자동 계산 · 105℃ 적용 시작 = P2 입력")
             + " · 노란 칸 = 미확인 조건 · 하이닉스 = 22,720÷12 (표시 반올림)")
        ws.freeze_panes = "C5"
        ws.print_area = f"B4:O{r}"
