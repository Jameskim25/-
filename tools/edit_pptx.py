# -*- coding: utf-8 -*-
"""첨부 PPT(SK trichem 양식) 수정 스크립트 — 기존 도형·표·배치를 재사용하고 수치는 재계산된 Excel에서 읽는다.

사용: python edit_pptx.py <base.pptx(구조 정리본)> <src2.pptx(원본 수정본)> <model.xlsx> <out.pptx>
base.pptx: 원본 1·2·3장 + 2027장 + 2027장 복제(2028용), 원본 2026장은 3장에 통합하기 위해 제거된 상태.
"""
import copy
import json
import sys

import openpyxl
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

BASE, SRC2, XLSX, OUT = sys.argv[1:5]
MAP = json.load(open(XLSX + ".map.json"))
wb = openpyxl.load_workbook(XLSX, data_only=True)
S = wb["06_Report_Summary"]; M = wb["04_Monthly_2026_2028"]; B = wb["02_Batch_Raw"]; I = wb["01_Inputs"]
M0 = MAP["M0"]


def K(k):
    return S.cell(row=MAP["KPI"][k], column=2).value


def f0(x):
    return f"{x:,.0f}"


def f1(x):
    return f"{x:,.1f}"


def sg(x, d=0):
    return f"{x:+,.{d}f}".replace("-", "−")


FONT = "Noto Sans KR"
E = 914400
prs = Presentation(BASE)
s1, s2, s3, s4, s5 = prs.slides


# ------------------------------------------------------------ helpers
def shapes(slide):
    return {sh.name: sh for sh in slide.shapes}


def set_par(p, texts):
    """paragraph의 기존 run 서식을 유지하며 텍스트 교체. texts: str 또는 run 텍스트 리스트."""
    if isinstance(texts, str):
        texts = [texts]
    runs = p.runs
    if not runs:
        r = p.add_run(); r.text = texts[0]
        runs = p.runs
    for i, t in enumerate(texts):
        if i < len(runs):
            runs[i].text = t
        else:
            last = p.runs[-1]._r
            nr = copy.deepcopy(last); last.addnext(nr)
            p.runs[-1].text = t
    for r in p.runs[len(texts):]:
        r._r.getparent().remove(r._r)


def set_lines(sh, lines):
    tf = sh.text_frame
    paras = tf.paragraphs
    while len(paras) < len(lines):
        last = paras[-1]._p
        last.addnext(copy.deepcopy(last))
        paras = tf.paragraphs
    for i, ln in enumerate(lines):
        set_par(paras[i], ln)
    for p in tf.paragraphs[len(lines):]:
        p._p.getparent().remove(p._p)


def pos(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = Emu(int(x * E))
    if y is not None: sh.top = Emu(int(y * E))
    if w is not None: sh.width = Emu(int(max(w, 0.005) * E))
    if h is not None: sh.height = Emu(int(h * E))


def delete(sh):
    sh._element.getparent().remove(sh._element)


def clone(slide, sh):
    el = copy.deepcopy(sh._element)
    slide.shapes._spTree.append(el)
    new = slide.shapes[-1]
    ids = [s.shape_id for s in slide.shapes]
    el.xpath("./*[1]/p:cNvPr")[0].set("id", str(max(ids) + 1))
    return new


def copy_fill(src, dst):
    """src 도형의 채우기·선·글자색을 dst에 복사 (spPr fill/ln, 첫 run 색)."""
    sp_s = src._element.spPr; sp_d = dst._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:ln"):
        for el in sp_d.findall(qn(tag)):
            sp_d.remove(el)
    geom = sp_d.find(qn("a:prstGeom"))
    anchor = geom if geom is not None else sp_d[-1]
    for tag in ("a:ln", "a:solidFill", "a:noFill"):
        el = sp_s.find(qn(tag))
        if el is not None:
            anchor.addnext(copy.deepcopy(el))
    if src.has_text_frame and dst.has_text_frame and src.text_frame.paragraphs[0].runs:
        sr = src.text_frame.paragraphs[0].runs[0]
        for p in dst.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = sr.font.size; r.font.bold = sr.font.bold
                if sr.font.color and sr.font.color.type == 1:
                    r.font.color.rgb = sr.font.color.rgb


def fill(sh, hexcol):
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(hexcol)


def runfmt(r, size, bold=False, color="3F3F3F"):
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color); r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def box(slide, x, y, w, h, color, text="", tsize=5.5, tcolor="FFFFFF", bold=True, line=None, dash=False, align="c"):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(max(w, 0.01)), Inches(h))
    sh.shadow.inherit = False
    if color:
        fill(sh, color)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    tf = sh.text_frame; tf.margin_left = tf.margin_right = Inches(0.01); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = {"c": PP_ALIGN.CENTER, "l": PP_ALIGN.LEFT}[align]
    if text:
        r = p.add_run(); r.text = text; runfmt(r, tsize, bold, tcolor)
    return sh


def tbox(slide, x, y, w, h, runs, align="l", wrap=True):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if isinstance(runs[0], list) else [runs]
    for i, pr in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"c": PP_ALIGN.CENTER, "l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT}[align]
        for t, size, bold, color in pr:
            r = p.add_run(); r.text = t; runfmt(r, size, bold, color)
    return sh


def cell_set(cell, text, tmpl=None):
    tf = cell.text_frame
    p = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if not p.runs:
        r = p.add_run()
        if tmpl is not None and tmpl.text_frame.paragraphs[0].runs:
            r._r.insert(0, copy.deepcopy(tmpl.text_frame.paragraphs[0].runs[0]._r.find(qn("a:rPr"))))
    set_par(p, text)


def table_rows(tbl, rows, tmpl_col=1):
    """rows: list of lists (len = ncols). None = 유지."""
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is None:
                continue
            c = tbl.cell(i, j)
            tmpl = tbl.cell(i, tmpl_col) if j != tmpl_col else tbl.cell(i, 0)
            cell_set(c, v, tmpl)


def unmerge_all(tbl):
    for i in range(len(tbl.rows)):
        for j in range(len(tbl.columns)):
            c = tbl.cell(i, j)
            if c.is_merge_origin:
                c.split()


def normalize_table(tbl, fills=None, mar=27000):
    """분할·병합으로 서식이 빠진 셀에 같은 행 기준 셀(1열)의 tcPr을 복사하고 여백·정렬·문단 끝 크기를 통일."""
    fills = fills or {}
    for i in range(len(tbl.rows)):
        ref = tbl.cell(i, 1)._tc.get_or_add_tcPr()
        for j in range(len(tbl.columns)):
            tc = tbl.cell(i, j)._tc
            tcPr = tc.get_or_add_tcPr()
            if tcPr.find(qn("a:lnB")) is None and j != 1:
                new = copy.deepcopy(ref)
                tc.replace(tcPr, new); tcPr = new
            tcPr.set("anchor", "ctr"); tcPr.set("marL", str(mar)); tcPr.set("marR", str(mar)); tcPr.set("marT", "0"); tcPr.set("marB", "0")
            if (i, j) in fills:
                for f in tcPr.findall(qn("a:solidFill")):
                    tcPr.remove(f)
                sf = etree.SubElement(tcPr, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", fills[(i, j)])
            for para in tbl.cell(i, j).text_frame.paragraphs:
                if j > 0:
                    para.alignment = PP_ALIGN.CENTER
                sz = para.runs[0].font.size if para.runs else None
                end = para._p.find(qn("a:endParaRPr"))
                if end is None:
                    end = etree.SubElement(para._p, qn("a:endParaRPr"))
                end.set("sz", str(int(sz.pt * 100)) if sz else "600")


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def mrow(col, y, rng=range(12)):
    return [M[f"{col}{M0 + (y - 2026) * 12 + k}"].value for k in rng]


# ============================================================ SLIDE 1 — As-is
d = shapes(s1)
set_lines(d["TextBox 5"], ["공용 정제기 1대 · 정제 45 h / 충진 전 57 h · 검사·충진은 다음 Batch 정제와 병렬"])
set_lines(d["TextBox 182"], ["공용 정제기 1대 · 103℃", "4개 고객 공동 (한솔 '27~)"])
tag_ok = d["Rounded Rectangle 197"]                     # IQC '2 h' (제공 시간 스타일)
for nm, txt in (("Rounded Rectangle 198", "2 h"), ("Rounded Rectangle 199", "정제 45 h/Batch"), ("Rounded Rectangle 200", "PQC + 제품 이송 합계 6 h")):
    copy_fill(tag_ok, d[nm]); set_lines(d[nm], [txt])
pos(d["Rounded Rectangle 200"], x=5.06, w=1.70)
delete(d["Rounded Rectangle 201"])
set_lines(d["TextBox 212"], ["5 Gal 글로브 박스 충진", "20 kg/병 · 필터 포함 · 하이닉스·CXMT", "2 h/병 · 약 9병 18 h (≈2.25 근무일)"])
set_lines(d["TextBox 216"], ["200 L 수동 충진", "이지켐 140 · 한솔 150 kg/용기 · 필터 포함", "이지켐 8 h/용기 · 한솔 시간 확인"])
set_lines(d["TextBox 226"], [["출하 ", "■", " 이지켐  ", "■", " 한솔('27~)"]])
d["TextBox 226"].text_frame.paragraphs[0].runs[3].font.color.rgb = RGBColor.from_string("7030A0")
set_lines(d["TextBox 227"], ["OQC·출하 : 5 Gal 9병 2 h · 200 L 1용기 2 h (운송 별도)"]); pos(d["TextBox 227"], x=3.90, w=3.10)
set_lines(d["TextBox 242"], ["정제기 1대 — 4개 고객 공동 (전용 설비 없음)"])
set_lines(d["TextBox 246"], ["103℃ · 정제 45 h · 충진 전 57 h · 190 kg/Batch"])
set_lines(d["TextBox 248"], ["5 Gal 20 kg/병 — 글로브 박스"])
set_lines(d["TextBox 250"], ["200 L 이지켐 140 · 한솔 150 kg — 수동"])
# ① Lead time bars (scale = 기존 57 h 막대 4.09 in)
set_lines(d["TextBox 258"], [["① 한 Batch 확인 공정시간 — 단순 합계", "   달력 납기 · 다음 Batch 투입 간격 아님"]])
k = 4.09 / 57
x0 = 1.79
set_lines(d["TextBox 269"], ["5 Gal 약 9병"]); set_lines(d["TextBox 277"], ["이지켐 200 L 1용기"])
for row, (iqc, ref, fqc, fil, oqc, unk, tot, fill_h, label) in enumerate((
        ("Rectangle 270", "Rectangle 271", "Rectangle 272", "Rectangle 273", "Rectangle 274", "Rectangle 275", "TextBox 276", 18, "충진 18 h"),
        ("Rectangle 278", "Rectangle 279", "Rectangle 280", "Rectangle 281", "Rectangle 282", "Rectangle 283", "TextBox 284", 8, "충진 8 h"))):
    pos(d[iqc], x=x0, w=2 * k)
    prep = clone(s1, d[ref]); fill(prep, "F4B183"); set_lines(prep, [""]); pos(prep, x=x0 + 2 * k, w=2 * k)
    pos(d[ref], x=x0 + 4 * k, w=45 * k); set_lines(d[ref], ["정제 45 h"])
    pq = clone(s1, d[fqc]); fill(pq, "9DC3E6"); pos(pq, x=x0 + 49 * k, w=6 * k)
    pqt = tbox(s1, x0 + 49 * k, d[ref].top / E + 0.04, 6 * k, 0.15, [("PQC·이송 6", 4.8, True, "1F3864")], align="c")
    pos(d[fqc], x=x0 + 55 * k, w=2 * k)
    pos(d[fil], x=x0 + 57 * k, w=fill_h * k); set_lines(d[fil], [label])
    pos(d[oqc], x=x0 + (57 + fill_h) * k, w=2 * k)
    pos(d[unk], x=x0 + (59 + fill_h) * k + 0.01, w=0.55)
    pos(d[tot], x=x0 + (59 + fill_h) * k + 0.62)
    v = K("lt_5g") if row == 0 else K("lt_ez")
    set_lines(d[tot], [[f"{v:.0f} h", " + 미확인"]] if len(d[tot].text_frame.paragraphs[0].runs) > 1 else [f"{v:.0f} h + 미확인"])
# 57 h 표시 (FQC 끝)
tbox(s1, x0 + 57 * k - 0.45, 3.93, 0.9, 0.12, [("▼ 충진 전 57 h", 5.5, True, "C00000")], align="c")
set_lines(d["TextBox 285"], ["※ IQC 2 + 준비·투입 2 + 정제 45 + PQC·이송 6 + FQC 2 = 57 h (중복 가산 없음) · OQC 2 h · 77 h·67 h는 확인시간 단순 합계 — 실제 달력 납기 아님 · 미확인: 대기·세척·전환·한솔 충진"])
# ② 11 Batch timetable — 바 길이를 시간 비율로 재배치 (1일 0.28 in = 24 h)
set_lines(d["TextBox 287"], [["② 11월 운전 예시 — 정제기 1대 월 11 Batch (과거 최대)", "   간격 65.5 h = 720÷11 (구성 확인)"]])
kd = 0.28 / 24
gx = 1.28
step = 720 / 11
for i in range(11):
    s = i * step
    bar = d[f"Rectangle {395 + 4 * i}"]; sm = d[f"Rectangle {396 + 4 * i}"]; gr = d[f"Rectangle {397 + 4 * i}"]
    pos(bar, x=gx + (s + 2) * kd, w=47 * kd)
    pos(sm, x=gx + (s + 49) * kd, w=6 * kd); fill(sm, "9DC3E6")
    g0 = s + 55; g1 = min(s + 77, 720)
    pos(gr, x=gx + g0 * kd, w=(g1 - g0) * kd)
set_lines(d["TextBox 440"], ["준비·정제 47 h"]); set_lines(d["TextBox 442"], ["PQC·이송 6 h"]); fill(d["Rectangle 441"], "9DC3E6")
set_lines(d["TextBox 444"], ["FQC·충진·OQC 22 h"])
pos(d["Rectangle 439"], x=6.55); pos(d["TextBox 440"], x=6.73, w=0.85); pos(d["Rectangle 441"], x=7.58); pos(d["TextBox 442"], x=7.76, w=0.75)
pos(d["Rectangle 443"], x=8.52); pos(d["TextBox 444"], x=8.70, w=1.0)
set_lines(d["TextBox 446"], [["과거 최대  ", f"11 Batch · {f0(K('mx_kgm'))} kg/월"], f"매월 11회 반복 참고 {K('mx_kgy12')/1000:.2f}톤/년 (11×190×12)"])
set_lines(d["TextBox 448"], [["시간 비교  ", f"{K('mx_ref'):.0f} h + {K('mx_gap'):.0f} h = 720 h"], "정제 11×45 + 225 h (점유·전환·대기·보수·월 경계 확인)"])
set_lines(d["TextBox 450"], [["이론 (정제만)  ", f"720÷45 = {K('th_bm'):.0f} Batch"], f"{f0(K('th_kgm'))} kg/월 · 실제 Capa. 아님 · 관측 {K('ob_int'):.1f} h → {K('ob_b30'):.1f}"])
set_lines(d["TextBox 452"], [["현재 연간 Capa.  ", "26.2톤 (기준값)"], f"26,200÷190 = {K('cp_b'):.1f} · 138×190 = {f0(K('cp_kground'))} kg · 월 {K('cp_bm'):.2f}"])
notes(s1, f"""[1장 As-is — 공정시간과 현재 생산능력 (Excel 01_Inputs B·C, 03_Capacity_Model 1~7절)]
■ 운영 전제: 정제기 1대를 하이닉스·CXMT·이지켐(·한솔 '27~)이 공동 사용 — 고객 전용 설비 없음. 앞 Batch 검사·충진 중 다음 Batch 정제 가능(병렬).
■ 공정시간 (사용자가 정한 자료 작성 기준 / 57시간에서 기타 공정시간 차감): IQC 2 + 준비·투입 2 + 정제 45 + PQC·제품 이송 합계 6 + FQC 2 = 충진 전까지 57 h. PQC와 이송은 그림에서 구분하되 시간은 합계 6 h(각 6 h 아님). 생산팀 대화 중 49 h 언급은 참고로만 보존.
■ 충진·OQC: 5 Gal 2 h/병, 약 9병 18 h(18 ÷ 8 h = 2.25 근무일, 현장 '약 2일' 설명) · 5 Gal OQC 9병 합계 2 h · 이지켐 200 L 수동 8 h/용기 · 200 L OQC 1용기 2 h · 한솔 수동·ARS 시간은 미확인(이지켐 8 h 복제 안 함). 충진·포장 1~1.5일은 개선 과제(달성 실적 아님).
■ 확인시간 단순 합계: 5 Gal 약 9병 57+18+2 = {K('lt_5g'):.0f} h · 이지켐 1용기 57+8+2 = {K('lt_ez'):.0f} h — 실제 달력 납기 아님, 다음 Batch 투입 간격 아님.
■ 정제기 점유 확인분 = 준비 2 + 정제 45 = {K('occ'):.0f} h (이송 중 점유분 미확인 → 47~53 h). 과거 일정 관측 평균 간격 {K('ob_int'):.2f} h → 비점유 약 {K('ob_vs47'):.1f} h(구성 미확인).
■ 이론(정제만): 720 ÷ 45 = {K('th_bm'):.0f} Batch/월 × 190 = {f0(K('th_kgm'))} kg/월 — 준비·전환·이송·세척·대기·보수 제외 비교값, 실제 Capa. 아님.
■ 과거 최대 11 Batch: 11×45 = {K('mx_ref'):.0f} h, 720 − 495 = {K('mx_gap'):.0f} h(전부 손실·충진으로 분류 안 함), 11×190 = 2,090 kg/월, 매월 11회 반복 참고 25.08톤/년. 11월 예시 막대는 65.5 h 균등 배분(준비·정제 47 h + PQC·이송 6 h + FQC·충진·OQC 22 h는 다음 Batch와 병렬).
■ 26.2톤: 26,200 ÷ 190 = {K('cp_b'):.3f} Batch 상당 → 138 × 190 = 26,220 kg = 26.22톤 → 26.2톤. 시간·가동률·수율을 조정해 맞추지 않음.""")

# ============================================================ SLIDE 2 — To-be
d = shapes(s2)
set_lines(d["TextBox 2"], ["To-be | 정제기 2대 운영과 투자 계획"])
set_lines(d["TextBox 5"], ["정제기 1·2 시간차 운전 · 공용 후공정 · '27.7 생산 시작 · 47.2톤/년 (연간 환산)"])
# 정제기 아이콘 2개 (기존 아이콘 축소·복제)
vessel = [sh for sh in s2.shapes if sh.name.startswith("Round Same Side Corner")][0]
noz, vtxt = d["Rectangle 177"], d["TextBox 179"]
for sh in (vessel, noz, vtxt):
    sh.left = Emu(int(sh.left - 0.36 * E))
set_lines(vtxt, ["정제기1"]); vtxt.text_frame.paragraphs[0].runs[0].font.size = Pt(5.5)
for sh in (vessel, noz, vtxt):
    c = clone(s2, sh); c.left = Emu(int(sh.left + 0.70 * E))
    if c.has_text_frame and c.text_frame.text:
        set_lines(c, ["정제기2"])
tbox(s2, 3.93, 1.40, 0.30, 0.14, [("시차", 5.5, True, "C00000")], align="c")
pos(d["Connector 193"], w=3.48 - 2.86)
c194 = d["Connector 194"]; old_end = (c194.left + c194.width) / E; pos(c194, x=4.60, w=old_end - 4.60)
set_lines(d["TextBox 182"], ["정제기 1·2 · 시간차 병행", "4개 고객 공용 배정 (전용 없음)"])
tag_ok = d["Rounded Rectangle 197"]
for nm, txt in (("Rounded Rectangle 198", "2 h"), ("Rounded Rectangle 199", "정제 45 h (설비별)"), ("Rounded Rectangle 200", "PQC + 제품 이송 합계 6 h")):
    copy_fill(tag_ok, d[nm]); set_lines(d[nm], [txt])
pos(d["Rounded Rectangle 199"], x=3.43, w=1.24); pos(d["Rounded Rectangle 200"], x=5.06, w=1.70)
delete(d["Rounded Rectangle 201"])
set_lines(d["TextBox 212"], ["5 Gal 글로브 박스 충진 (유지)", "20 kg/병 · 필터 포함 · ARS 미적용", "2 h/병 · 약 9병 18 h"])
set_lines(d["TextBox 216"], ["200 L 수동 → ARS 자동 충진", "이지켐 140 kg (전환 월 확인) · 한솔 150 kg", "한솔 '27.1~6 수동 · '27.7~ ARS · 시간 확인"])
set_lines(d["TextBox 226"], [["출하 ", "■", " 이지켐  ", "■", " 한솔"]])
d["TextBox 226"].text_frame.paragraphs[0].runs[3].font.color.rgb = RGBColor.from_string("7030A0")
set_lines(d["TextBox 227"], ["OQC·출하 : 5 Gal 9병 2 h · 200 L 1용기 2 h (운송 별도)"]); pos(d["TextBox 227"], x=3.90, w=3.10)
set_lines(d["TextBox 240"], ["개선 영향 구간"])
set_lines(d["Rounded Rectangle 241"], ["정제기 2대"]); d["Rounded Rectangle 241"].text_frame.paragraphs[0].runs[0].font.size = Pt(6)
set_lines(d["TextBox 242"], ["정제기 1·2 시간차 병행 · '27.7 생산", "+21톤/년 (12개월 환산) · 단순 2배 아님"])
set_lines(d["TextBox 244"], ["200 L · '26 설치 → 한솔 '27.7~", "이지켐 전환 월 확인 · Capa. 가산 없음"])
set_lines(d["TextBox 246"], ["5 Gal 글로브 박스 · 2 h/병", "하이닉스 · CXMT · ARS 미적용"])
set_lines(d["TextBox 247"], ["· 정제 45 h → 22.5 h 단축 아님 (설비 2대)", "· 공용 후공정(검사·Tank·충진) 병목 검토"])
# ① 일정 표
set_lines(d["TextBox 255"], [["① 투자 · 운영 일정", "   '27.6까지 정제기 1대 · 26.2톤/년 → '27.7부터 정제기 2대 · 47.2톤/년 (연간 환산)"]])
tbl = d["Table 256"].table
cell_set(tbl.cell(2, 2), "공사 (2~5월) · 공사 중 가동 여부 확인")
cell_set(tbl.cell(3, 6), "시운전")
cell_set(tbl.cell(4, 0), "■ 정제기 2대 적용 생산")
cell_set(tbl.cell(4, 7), "정제기 2대 생산 시작 '27.7 · 최초 12개월 '27.7~'28.6 (47.2톤 환산)")
cell_set(tbl.cell(5, 0), "■ ARS 200 L 충진")
cell_set(tbl.cell(5, 1), "'26 설치 → '27 운영 계획 · 한솔 1~6월 수동 → 7월~ ARS · 이지켐 전환 월 확인 · 5 Gal 미적용")
# ② 2대 시차 운전 Gantt — 기존 11월 예시(1대 21 Batch) 삭제 후 같은 영역에 작성
for sh in list(s2.shapes):
    t = sh.top / E
    if (5.25 <= t < 6.245 and sh.name not in ("TextBox 258",)) or (4.95 <= t < 5.2 and sh.left / E > 6.0):
        delete(sh)
off, gint = K("off"), K("g_int")
set_lines(d["TextBox 258"], [["② 정제기 2대 시간차 운전 개념", f"   운영 개념 설명용 가정: 기동 시차 {off:.0f} h · 설비별 간격 {gint:.1f} h (47.2톤 역산 참고) — 미확정"]])
X0, X1, H = 1.28, 9.69, 216
kk = (X1 - X0) / H
for hh in range(0, 217, 24):
    tbox(s2, X0 + hh * kk - 0.18, 5.26, 0.36, 0.11, [(f"{hh}" + (" h" if hh == 216 else ""), 5.3, False, "7F7F7F")], align="c")
    box(s2, X0 + hh * kk, 5.38, 0.004, 0.84, "E7E6E6")
lanes = [("정제기 1", 5.40), ("정제기 2", 5.61), ("PQC·이송·FQC", 5.82), ("5 Gal 충진 (공용)", 6.03)]
for lab, yy in lanes:
    tbox(s2, 0.31, yy + 0.02, 0.95, 0.16, [(lab, 6.3, True, "1A1A1A")])


def gbar(lane_y, a, b, color, text="", tcolor="FFFFFF", **kw):
    a = max(a, 0); b = min(b, H)
    if b <= a:
        return
    box(s2, X0 + a * kk, lane_y, (b - a) * kk, 0.17, color, text if (b - a) * kk > 0.55 else "", 5.2, tcolor, **kw)


for u, (lab, yy) in enumerate(lanes[:2]):
    o = 0 if u == 0 else off
    for b in range(4):
        s = o + b * gint
        gbar(yy, s + 2, s + 4, "F4B183")
        gbar(yy, s + 4, s + 49, "FF7900", f"정제기{u+1} B{b+1} 정제 45 h")
        gbar(yy, s + 49, s + gint + 2, None, "기타 (Mix·이송·대기 확인)", "7F6000", line="BF9000", dash=True)
        gbar(5.82, s + 49, s + 57, "2E75B6")
        gbar(6.03, s + 57, s + 75, "548235", f"{u+1}-B{b+1} 18 h")
tbox(s2, 0.31, 6.21, 9.4, 0.06, [("", 3, False, "FFFFFF")])
set_lines(d["TextBox 454"], [f"Capa. (연간 환산)", "26.2 → 47.2톤/년", "'27.7~'28.6 12개월 · 2027 실제 생산량 아님"])
set_lines(d["TextBox 456"], ["단순 2배 vs 개선 환산", "52.4 vs 47.2톤", f"차이 {K('x2_gap'):.1f}톤 = Mix·리사이클 준비 등 (확정 손실 아님)"])
set_lines(d["TextBox 458"], ["설비별 등가 투입 간격", f"약 {K('rf_int_each'):.1f} h (2대)", f"47.2톤÷190 역산 · 현재 1대 {K('cp_int'):.1f} h · 45 h 단축 아님"])
set_lines(d["TextBox 460"], ["105℃ 검토안", "50.2톤 (미확정)", "+3톤 · 열 안정성·Dimer·Unknown·Yield · 승인·시점 미정"])
notes(s2, f"""[2장 To-be — 정제기 2대 운영과 투자 계획 (Excel 03_Capacity_Model 8절, 05_Reflux_Scenarios)]
■ 정제기: As-is 1대 단일 설비 순차 생산 → To-be 정제기 1·2를 시간차를 두고 병행 운전. 정제기 1을 먼저 투입·가동한 뒤 투입·초류 진행 상황과 충진 처리능력을 고려해 정제기 2를 기동하고, 두 설비 운전 구간을 겹치되 제품이 한꺼번에 충진으로 몰리지 않도록 배치. 고객별 전용 정제기 없음 — 전체 설비 자원 안에서 고객 물량 배정.
■ 시차·설비별 간격: 미확정. Gantt의 시차 {off:.0f} h·설비별 간격 {gint:.1f} h는 운영 개념 설명용 가정(동시 기동·24 h 고정 아님). 설비별 간격 {gint:.1f} h = 47.2톤 ÷ 190 kg = {47200/190:.1f} Batch/년을 2대로 나눈 역산 참고값(현재 1대 26.2톤 등가 {K('cp_int'):.1f} h 대비 +{K('rf_int_diff'):.1f} h — Mix·리사이클 준비 등 가능성, 구성 미확인). 정제 45 h를 22.5 h로 단축하는 계산이 아님.
■ Capa.: 26.2 + 21 = 47.2톤/년(최초 12개월 '27.7~'28.6 연간 환산). 현재 단순 2배 52.4톤과의 차이 {K('x2_gap'):.1f}톤은 보조 작업·공유 자원 등 가능성 — 확정 손실로 표시하지 않음. 2027 실제 생산량은 공사 중 가동·시운전 양품·출하 승인·초기 안정화를 반영해 별도 산정(4장).
■ 충진: 5 Gal은 글로브 박스 유지(ARS 미적용). 200 L ARS는 '26 설치 → '27 운영 계획 — 한솔 '27.1~6 수동, '27.7~ ARS, 2028 전월 ARS. 이지켐 ARS 전환 월은 근거('27~ 운영 계획)만 있어 확인 필요(한솔 7월 전환 자동 적용 안 함). ARS 효과는 정제 Capa.에 가산하지 않음.
■ 공용 후공정: 2대 합산 평균 투입 간격 {K('rf_int_all'):.1f} h — 1교대(8 h/일) 충진 시 18 h 작업 = 경과 54 h로 Tank·충진 회전 부족({K('tk1_tobe'):.1f} h), 2교대 시 여유 {K('tk2_tobe'):.1f} h (Product Tank 1기 가정, 03 시트 10절).
■ 2028 검토안: 103→105℃, +3톤 → 50.2톤(산술) · 목표 2028년 이후 50톤 수준 · 미확정·적용 시점 미정·검증 및 승인 필요(열 안정성, Dimer, Unknown impurity, Yield).""")

# ============================================================ SLIDE 3 — 과거 Batch + 2026
d = shapes(s3)
set_lines(d["TextBox 2"], ["과거 Batch 분석 및 2026년 기준 검토"])
set_lines(d["TextBox 5"], [f"관측 간격 약 {K('ob_int'):.1f} h (6/2 #39~10/29 #92) · 2026 출하 19.22톤 vs 현재 26.2톤"])
set_lines(d["TextBox 244"], ["② 표시일자 간격 분포 (#39~#92 · 54개 일자 · 53간격)"])
set_lines(d["TextBox 261"], [["평균 ", f"{K('ob_days')/K('ob_n'):.2f}일 = 약 {K('ob_int'):.1f} h", f"  ({K('ob_lo'):.1f}~{K('ob_hi'):.1f} h)"],
                             f"중앙값 {K('ob_med'):.0f} h (3일) · 월 경계 포함 · 5월 일부 제외 · 날짜만 있는 계획"])
# ③ 표
set_lines(d["TextBox 263"], [["③ 월 생산능력 비교", "   정제기 1대 · 월 720 h · 연 환산은 참고값"]])
tbl = d["Table 264"].table
Cm = wb["03_Capacity_Model"]; c0 = MAP["CMP0"]
cm = lambda r, c: Cm.cell(row=c0 + r, column=c).value
table_rows(tbl, [[None] * 6,
                 ["이론 (정제 45 h만)", "45.0", f"{cm(0,3):.0f}", f0(cm(0,4)), f"{cm(0,5):.1f} (참고)", "720÷45 · 준비·전환·이송·세척·대기·보수 제외"],
                 ["과거 월 최대 11 Batch", f"{cm(1,2):.1f}", "11", f0(cm(1,4)), f"{cm(1,5):.2f}", "정제 11×45 = 495 h + 225 h (구성 확인)"],
                 ["과거 일정 추정 (6~10월)", f"약 {cm(2,2):.1f}", f"약 {cm(2,3):.1f}", f"약 {f0(cm(2,4))}", f"약 {cm(2,5):.1f}", "149일×24÷53 · 중앙값 72 h · 표시일 기준"],
                 ["현재 Capa. 26.2톤 환산", f"{cm(3,2):.1f}", f"{K('cp_bm'):.2f}", f0(26200 / 12), "26.2", "26,200÷190 = 137.9 Batch 상당 · 기준값"]])
gf = d["Table 264"]; pos(gf, y=3.10, h=5 * 0.175)
for r_ in tbl.rows:
    r_.height = Emu(int(0.175 * E))
# ④ 병렬 Gantt 압축·재배치
ymap = {4.76: 4.47, 5.07: 4.66, 5.39: 4.85, 5.71: 5.04}
pos(d["Rectangle 265"], y=4.02); pos(d["TextBox 266"], y=3.99)
set_lines(d["TextBox 266"], [f"④ 정제 · 검사 · 충진 병렬 운영 (관측 간격 {K('ob_int'):.1f} h 기준 · 정제 45 h)"])
for sh in s3.shapes:
    t = round(sh.top / E, 2)
    if sh.name.startswith("TextBox") and abs(t - 4.57) < 0.02:
        pos(sh, y=4.27)
    if sh.name.startswith("Connector") and abs(t - 4.74) < 0.02:
        pos(sh, y=4.40, h=0.82)
for nm in ("TextBox 283", "TextBox 284", "TextBox 285", "TextBox 286"):
    sh = d[nm]; t = round(sh.top / E, 2)
    near = min(ymap, key=lambda v: abs(v - t))
    pos(sh, y=ymap[near] - 0.02, h=0.18)
set_lines(d["TextBox 283"], ["정제기 (1대)"]); set_lines(d["TextBox 284"], ["PQC·이송·FQC"])
kx = (2.16 - 1.46) / 30
g = K("ob_int")
segs = {  # name: (lane_y, start, end, text)
    "Rectangle 287": (4.47, 0, 47, "B n 준비·정제 47 h"), "Rectangle 288": (4.47, 47, g, ""),
    "Rectangle 289": (4.66, 47, 55, ""), "Rectangle 290": (4.85, 55, 73, ""), "Rectangle 291": (5.04, 73, 75, ""),
    "Rectangle 292": (4.47, g, g + 47, "B n+1 준비·정제 47 h"), "Rectangle 293": (4.47, g + 47, 2 * g, ""),
    "Rectangle 294": (4.66, g + 47, g + 55, ""), "Rectangle 295": (4.85, g + 55, g + 73, ""), "Rectangle 296": (5.04, g + 73, g + 75, ""),
    "Rectangle 297": (4.47, 2 * g, 2 * g + 47, "B n+2 준비·정제 47 h"), "Rectangle 298": (4.66, 2 * g + 47, 2 * g + 55, ""),
    "Rectangle 299": (4.85, 2 * g + 55, min(2 * g + 73, 215), "")}
for nm, (yy, a, b, txt) in segs.items():
    sh = d[nm]; pos(sh, x=1.46 + a * kx, y=yy, w=(b - a) * kx, h=0.16)
    if txt:
        set_lines(sh, [txt])
fill(d["Rectangle 289"], "2E75B6"); fill(d["Rectangle 294"], "2E75B6"); fill(d["Rectangle 298"], "2E75B6")
set_lines(d["TextBox 300"], [f"← 기타 약 {g-47:.1f} h (구성 확인)"]); pos(d["TextBox 300"], x=1.46 + 55 * kx + 0.03, y=4.67, w=1.15)
set_lines(d["TextBox 301"], ["↑ 충진은 다음 Batch 정제와 병렬"]); pos(d["TextBox 301"], x=1.46 + 75 * kx + 0.05, y=5.05, w=1.4)
set_lines(d["TextBox 302"], ["※ IQC(원료 Lot 공통 여부)·이송 중 정제기 점유·충진 인력·글로브 박스 근무시간이 다음 Batch를 제한하는지 확인"])
pos(d["TextBox 302"], y=5.22)
pos(d["Rectangle 303"], y=4.02, h=1.36); pos(d["TextBox 304"], y=4.05, h=1.30)
set_lines(d["TextBox 304"], ["추정 근거 · 가정 · 한계",
                             f"데이터 : 6/2 #39~10/29 #92 ({K('ob_n_dates'):.0f}개 일자 · {K('ob_n'):.0f}간격) · 5월 일부 제외",
                             f"방법 : {K('ob_days'):.0f}일 × 24 ÷ {K('ob_n'):.0f} = {K('ob_int'):.1f} h (범위 {K('ob_lo'):.1f}~{K('ob_hi'):.1f}) · 중앙값 72 h",
                             "가정 : 표시일 = 매 Batch 같은 이벤트 · 번호 Batch 모두 수행",
                             "한계 : 날짜만 있는 계획 (실적 아님) · D일 = (D±1)×24 h, 2일≠48 h",
                             f"차이 : 정제 45 h 대비 +{K('ob_vs45'):.1f} h · 점유 47 h 대비 +{K('ob_vs47'):.1f} h (구성 확인)",
                             "확인 : 표시일 의미 · 실제 시각 · Total 46 vs 번호 54 · 색상"])
for i_, p_ in enumerate(d["TextBox 304"].text_frame.paragraphs):
    for r_ in p_.runs:
        r_.font.size = Pt(7.5 if i_ == 0 else 6.2)
# ⑤ 2026 요약 — 원본 2026장의 표를 복사해 사용
src2 = Presentation(SRC2)
t26 = [sh for sh in src2.slides[3].shapes if sh.has_table][0]
el = copy.deepcopy(t26._element)
s3.shapes._spTree.append(el)
gf26 = s3.shapes[-1]
el.xpath("./p:nvGraphicFramePr/p:cNvPr")[0].set("id", str(max(s.shape_id for s in s3.shapes) + 1))
mk = clone(s3, d["Rectangle 262"]); pos(mk, y=5.47)
hd = clone(s3, d["TextBox 263"]); pos(hd, y=5.44)
set_lines(hd, [["⑤ 2026 출하 제시분 vs 현재 Capa. (kg)", f"   19.22톤 (CXMT·한솔 미제시) = 26.2톤의 {K('sh26')/26.2*100:.1f}% · 차이 {K('d26_now'):.2f}톤 (가동률·판매 가능량 아님)"]])
t = gf26.table
trs = t._tbl.findall(qn("a:tr"))
for tr in trs[7:]:
    t._tbl.remove(tr)
plan_t = mrow("M", 2026); plan_n = mrow("L", 2026); cap26 = mrow("S", 2026)
pl = [("" if a in ("", None) else f"{int(a)*190:,}") for a in plan_t]
table_rows(t, [[None] * 14, [None] * 14, [None] * 14, [None] * 14,
               ["계획 Total×190 (번호 ①)"] + pl + ["8,740*"],
               ["추정 생산 가능 (관측)"] + [f0(v) for v in cap26] + [f0(K("p26"))],
               ["추정 − 출하"] + [sg(a - b) for a, b in zip(cap26, mrow("Y", 2026))] + [sg(K("p26") - 19220)]])
pos(gf26, x=0.31, y=5.68, w=9.37, h=7 * 0.158)
for r_ in t.rows:
    r_.height = Emu(int(0.158 * E))
for i in range(len(t.rows)):
    for j in range(len(t.columns)):
        for p in t.cell(i, j).text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(5.8)
        t.cell(i, j).margin_top = t.cell(i, j).margin_bottom = 0
normalize_table(t, mar=20000)
note26 = tbox(s3, 0.31, 6.89, 5.8, 0.1, [(f"* 6~10월 합계 (번호×190 기준 10,260) · 계획은 원자료 생산계획(실적 아님) · 추정 생산 가능 = 관측 {K('ob_int'):.1f} h·190 kg·달력 연속({K('pb26'):.0f} Batch) · 상세 Excel 02·04", 5.3, False, "7F7F7F")])
notes(s3, f"""[3장 과거 Batch 분석 및 2026년 기준 검토 (Excel 02_Batch_Raw, 03_Capacity_Model 3~7절, 04_Monthly_2026_2028)]
■ 원자료(첨부 PPT 원본·수정본 3장 발표자 노트, 원본 분포도 좌표, 사용자 지시 목록 3중 대조 → 차이 {K('chk_diff'):.0f}건). 원본 이미지 5개는 미첨부. 연도 미표기 → 2026 가정.
- 5월(일부): 5/23 #37, 5/26 #38
- 6월: 6/2 #39, 6/5 #40, 6/8 #41, 6/11 #42, 6/14 #43, 6/17 #44, 6/20 #45, 6/23 #46, 6/25 #47, 6/27 #48, 6/28 #49
- 7월: 7/1 #50, 7/3 #51, 7/6 #52, 7/8 #53, 7/11 #54, 7/14 #55, 7/17 #56, 7/20 #57, 7/23 #58, 7/26 #59, 7/29 #60
- 8월: 8/1 #61, 8/4 #62, 8/7 #63, 8/10 #64, 8/13 #65, 8/16 #66, 8/19 #67, 8/22 #68, 8/25 #69, 8/28 #70, 8/31 #71
- 9월: 9/3 #72, 9/5 #73, 9/7 #74, 9/9 #75, 9/11 #76, 9/14 #77, 9/19 #78, 9/23 #79, 9/26 #80, 9/28 #81
- 10월: 10/1 #82, 10/4 #83, 10/7 #84, 10/10 #85, 10/12 #86, 10/15 #87, 10/18 #88, 10/20 #89, 10/23 #90, 10/26 #91, 10/29 #92
■ 집계: 원자료 Total 6월 8 / 7월 11 / 8월 8 / 9월 8 / 10월 11 = 46, 번호 개수 11/11/11/10/11 = 54 — 번호 개수를 양품 완료 Batch로 확정하지 않음. 색상·표시일자 의미 확인.
■ 관측 간격: 주 범위 6/2 #39~10/29 #92, 54개 일자·53개 간격, {K('ob_days'):.0f}일 × 24 ÷ 53 = {K('ob_int'):.2f} h(범위 {K('ob_lo'):.2f}~{K('ob_hi'):.2f}), 중앙값 72 h, 3일 간격 {K('ob_share3')*100:.0f}%. 월 경계 간격 포함, 5월 일부 제외.
- 한계: 날짜만 있는 계획 — 실제 정제시간과 구분. 정제 45 h 대비 +{K('ob_vs45'):.1f} h, 점유 47 h 대비 +{K('ob_vs47'):.1f} h는 PQC·이송·세척·대기·계획 여유 등 가능성(구성 미확인).
- 57 h 연속 가정에서 확인 필요였던 고밀도 구간(6/23~7/3, 9/3~9/11, 4간격 평균 최대 54 h)은 점유 하한 47 h 기준으로는 양립 가능 — 이송 중 점유·세척 반영 시 재확인.
■ 비교: 이론 720÷45 = 16 Batch·3,040 kg/월(실제 Capa. 아님) · 과거 최대 11 Batch 495 h + 225 h · 관측 {K('ob_b30'):.1f} Batch/30일(31일 {K('ob_b31'):.1f}) · 26.2톤 = 월 11.49 Batch 상당(역산 {K('cp_int'):.1f} h).
■ 2026: 제시 출하 19,220 kg(하이닉스 16,980 + 이지켐 2,240) — CXMT·한솔 미제시(0으로 확정 안 함). 현재 26.2톤 대비 {K('sh26')/26.2*100:.1f}%, 산술 차이 6.98톤. 관측 기준 추정 생산 가능 {f0(K('p26'))} kg({K('pb26'):.0f} Batch, 정지·보수 미반영 상한) — 9~12월 월 출하 1,940~2,040 kg은 월 생산 가능량(10~11 Batch)에 근접, 재고 활용 필요. 상세 월별은 Excel 04 시트.""")

# ============================================================ SLIDE 4·5 — 2027 / 2028
def ensure_merge(tbl, r, c1, c2):
    tbl.cell(r, c1).merge(tbl.cell(r, c2))
    for j in range(c1 + 1, c2 + 1):
        cell_set(tbl.cell(r, j), "")


def year_slide(slide, yr):
    d = shapes(slide)
    tbl = d["Table 166"].table
    unmerge_all(tbl)
    hx, cx, ez, hs, tot = mrow("U", yr), mrow("V", yr), mrow("W", yr), mrow("X", yr), mrow("Y", yr)
    capr, need, gb = mrow("T", yr), mrow("Z", yr), mrow("AB", yr)
    cez, chs, mez, mhs = mrow("AD", yr), mrow("AE", yr), mrow("AF", yr), mrow("AG", yr)
    sr = MAP["SR"][str(yr)]
    tot_y = M[f"Y{sr}"].value
    ann = lambda col: M[f"{col}{sr}"].value
    rows = [[None] * 14]
    if yr == 2027:
        rows.append(["적용 조건", "현재 1대", "공사 (가동 여부 확인)", "", "", "", "시운전", "정제기 2대 · 초기 안정화 확인", "", "", "", "", "", ""])
        rows.append(["참고 Capa. 월 환산 (÷12)"] + [f1(v) for v in capr] + [f0(ann("T"))])
        rows.append(["실제 생산 가능량", f0(K("p27_1")), "미확인 (공사 중 가용시간)", "", "", "", "미확인", "미확인 (설비별 간격·Batch량)", "", "", "", "", "", f"부분 {f0(K('p27_1'))}"])
    else:
        rows.append(["적용 조건", "정제기 2대 · 최초 12개월 ('27.7~'28.6)", "", "", "", "", "", "정제기 2대 · 47.2톤 기준 (105℃ 적용 시점 미정)", "", "", "", "", "", ""])
        rows.append(["참고 Capa. 47.2 (÷12)"] + [f1(v) for v in capr] + [f0(ann("T"))])
        rows.append(["105℃ 검토안 50.2 (÷12)"] + [f1(50200 / 12)] * 12 + ["50,200"])
    rows.append(["참고 Capa. − 출하"] + [sg(a - b) for a, b in zip(capr, tot)] + [sg(ann("T") - tot_y)])
    rows.append(["■ SK하이닉스"] + [f1(v) for v in hx] + [f0(ann("U"))])
    rows.append(["■ CXMT"] + [f0(v) for v in cx] + [f0(ann("V"))])
    rows.append(["■ 이지켐"] + [f0(v) for v in ez] + [f0(ann("W"))])
    rows.append(["■ 한솔"] + [f0(v) for v in hs] + [f0(ann("X"))])
    rows.append(["출하 합계"] + [f1(v) for v in tot] + [f0(tot_y)])
    rows.append(["필요 Batch (÷190)"] + [f1(v) for v in need] + [f1(ann("Z"))])
    rows.append(["5 Gal 충진 h (병 상당×2)"] + [f1(v) for v in gb] + [f0(ann("AB"))])
    rows.append(["이지켐 용기 (140 kg)"] + [f0(v) for v in cez] + [f"{ann('AD'):.0f}용기"])
    rows.append(["한솔 용기 (150 kg)"] + [f0(v) for v in chs] + [f"{ann('AE'):.0f}용기"])
    if yr == 2027:
        rows.append(["200 L 충진 방식", "한솔 수동 · 이지켐 ARS 전환 월 확인 ('27~ 계획)", "", "", "", "", "", "한솔 ARS · 이지켐 ARS (충진시간 확인 입력)", "", "", "", "", "", "확인"])
        rows.append(["상·하반기 비교 (t)", f"상반기 출하 {K('sh27h1'):.2f} vs 26.2×6/12 = {K('cap27h1'):.1f} (참고)", "", "", "", "", "",
                     f"하반기 출하 {K('sh27h2'):.2f} vs 47.2×6/12 = {K('cap27h2'):.1f} → {abs(K('d27_h2')):.2f} 부족 방향", "", "", "", "", "", f"{K('sh27'):.1f}"])
        rows.append(["재고 (기초 / 기말)", "기초재고 미입력 → 계산 보류 · 2027.12 기말재고 → 2028.1 기초재고로 연결 (Excel 04)", "", "", "", "", "", "", "", "", "", "", "", "보류"])
    else:
        rows.append(["200 L 충진 방식", "이지켐 · 한솔 ARS (충진시간 확인 입력) · 5 Gal 글로브 박스 유지", "", "", "", "", "", "", "", "", "", "", "", "확인"])
        rows.append(["실제 생산 가능량", "미확인 — 설비별 간격·Batch량·양품률 입력 시 계산 (Excel 04) · 연간 환산값과 구분", "", "", "", "", "", "", "", "", "", "", "", "보류"])
        rows.append(["재고 (기초 / 기말)", "2027.12 기말재고 → 2028.1 기초재고 연결 · 미확인 시 0 초기화 안 함", "", "", "", "", "", "", "", "", "", "", "", "보류"])
    table_rows(tbl, rows)
    for r_ in tbl.rows:
        r_.height = Emu(int(0.168 * E))
    pos(d["Table 166"], h=len(tbl.rows) * 0.168)
    for i in range(len(tbl.rows)):
        for j in range(14):
            c = tbl.cell(i, j); c.margin_top = c.margin_bottom = 0
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(6.0)
    set_lines(d["TextBox 10"], [f"- {4 if yr == 2027 else 5} -"]); set_lines(d["타원 6"], [str(4 if yr == 2027 else 5)])
    if yr == 2027:
        merges = [(1, 2, 5), (1, 7, 12), (3, 2, 5), (3, 7, 12), (14, 1, 6), (14, 7, 12), (15, 1, 6), (15, 7, 12), (16, 1, 12)]
    else:
        merges = [(1, 1, 6), (1, 7, 12), (14, 1, 12), (15, 1, 12), (16, 1, 12)]
    for r, a, b in merges:
        ensure_merge(tbl, r, a, b)
    UNKF, CONF = "FFF2CC", "FBE5D6"
    fills = {(1, 2): CONF, (1, 6): UNKF} if yr == 2027 else {}
    if yr == 2027:
        fills.update({(3, 2): UNKF, (3, 6): UNKF, (3, 7): UNKF})
    else:
        fills.update({(15, 1): UNKF})
    normalize_table(tbl, fills)
    # 하단 블록(②·③)을 표 길이에 맞춰 이동, 주석은 표 바로 아래
    for sh in slide.shapes:
        t = sh.top / E
        if 4.45 <= t < 6.8:
            sh.top = Emu(int(sh.top + 0.21 * E))
    pos(d["TextBox 167"], y=4.50, h=0.2)
    # 음수 강조
    for i in range(len(tbl.rows)):
        for j in range(1, 14):
            for p in tbl.cell(i, j).text_frame.paragraphs:
                for r in p.runs:
                    if r.text.startswith("−"):
                        r.font.color.rgb = RGBColor.from_string("C00000")
    return d


d = year_slide(s4, 2027)
set_lines(d["TextBox 2"], ["2027년 월별 출하와 공급 대응"])
set_lines(d["TextBox 5"], [f"출하계획 43.4톤 · 현재 26.2톤 대비 17.2톤 초과 · 하반기 24.56 vs 참고 23.6톤"])
set_lines(d["TextBox 165"], [["① 2027 월별 출하 · 참고 Capa. · 충진 (kg)", "   참고 Capa. = 연간 환산 ÷ 12 (공사·시운전·안정화 미반영) · 실제 생산 가능량은 조건 확인 후 계산"]])
set_lines(d["TextBox 167"], [f"* 하이닉스 22,720÷12 = 1,893.33 kg/월(표시 반올림) · CXMT 7,000 · 이지켐 10,080 (72용기) · 한솔 3,600 (24용기) · 5 Gal 병 상당 = (하이닉스+CXMT)÷20, 실제 20 kg 정수 배분 · 실제 생산 1월 = 관측 {K('basis_int'):.1f} h 기준 {K('pb27_1'):.0f} Batch"])
set_lines(d["TextBox 169"], [["② 2027 연간 환산 기준 비교", "   확정 부족량·판매 가능량 아님"]] if len(d["TextBox 169"].text_frame.paragraphs[0].runs) > 1 else ["② 2027 연간 환산 기준 비교 (확정 부족량 아님)"])
kp4 = [("현재 Capa. 대비", ["43.4 vs 26.2", " 톤"], "17.2톤 초과 → 개선·선행 생산 필요"),
       ("개선 환산 47.2톤 대비", [f"+{K('d27_rf'):.1f}", " 톤"], "연간 환산값 — 2027 실제 여유 아님"),
       ("하반기 출하 vs 참고 Capa.", ["24.56 vs 23.6", " 톤"], f"{abs(K('d27_h2')):.2f}톤 부족 방향 (47.2×6/12 균등 참고)"),
       ("2027 생산 − 출하 (시나리오)", [f"{K('S2_d27')/1000:.1f} ~ {K('S3_d27')/1000:.1f}".replace("-", "−"), " 톤"], "S2 공사 중 유지 ~ S3 50%·안정화 3개월 (Excel 05)")]
for (a, b_, c_), base in zip(kp4, (171, 175, 179, 183)):
    set_lines(d[f"TextBox {base}"], [a]); set_lines(d[f"TextBox {base+1}"], [b_]); set_lines(d[f"TextBox {base+2}"], [c_])
set_lines(d["TextBox 187"], ["③ 정제 · 충진 자원 제약 (월 기준)"])
kp4b = [("정제기 1~6월 (1대)", ["필요 14.5~17.5 vs 약 11", " Batch/월"], f"관측 {K('ob_int'):.1f} h 기준 · 공사 중 가동 미확인"),
        ("정제기 7~12월 (2대)", ["필요 21.5 vs 20.7", " Batch/월"], "47.2톤÷12÷190 · 설비별 간격 미확인"),
        ("5 Gal 글로브 박스", ["189.3 ~ 267.3", " h/월"], f"연 {K('gb27'):,.0f} h · 1교대 176 h의 {K('gb_ld1')*100:.0f}% (최대)"),
        ("200 L (이지켐+한솔)", ["6 → 10", " 용기/월"], f"연 {K('c200_27'):.0f}용기 · OQC 2 h/용기 · 한솔·ARS 시간 확인")]
for (a, b_, c_), base in zip(kp4b, (189, 193, 197, 201)):
    set_lines(d[f"TextBox {base}"], [a]); set_lines(d[f"TextBox {base+1}"], [b_]); set_lines(d[f"TextBox {base+2}"], [c_])
notes(s4, f"""[4장 2027년 월별 출하와 공급 대응 (Excel 01_Inputs K, 04_Monthly_2026_2028, 05_Reflux_Scenarios)]
■ 출하계획(kg): 하이닉스 22,720÷12 = 1,893.33/월(표시 반올림, 연간 22,720 유지) · CXMT 1~2월 0 / 3~6월 580 / 7~12월 780 = 7,000 · 이지켐 1~6월 560 / 7~12월 1,120 = 10,080 · 한솔 1~12월 300 = 3,600 → 합계 43,400 kg = 43.4톤. 월 합계 1~2월 2,753.3 / 3~6월 3,333.3 / 7~12월 4,093.3.
■ 연간 환산 비교: 현재 26.2톤 대비 17.2톤 초과 · 개선 후 47.2톤과 3.8톤 차이(연간 환산 — 2027 실제 여유 아님). 상반기 출하 18.84톤 / 하반기 24.56톤, 하반기 균등 가동 참고 47.2×6/12 = 23.6톤 → 0.96톤 부족 방향.
■ 2027 실제 생산: 1월 현재 조건(관측 {K('basis_int'):.2f} h) {K('pb27_1'):.0f} Batch = {f0(K('p27_1'))} kg. 2~5월 공사 중 가동, 6월 시운전 양품·출하 승인, 7월 이후 설비별 간격·Batch량·초기 안정화가 미확인 → 연간 실제 생산량은 부분합계로 보류(0 또는 정상값 임의 입력 안 함).
■ 시나리오(참고, Excel 05): S1 공사 중 중단 → 2027 생산 {K('S1_p')/1000:.1f}톤(생산−출하 {K('S1_d27')/1000:.1f}톤) · S2 공사 중 현재 유지 → {K('S2_p')/1000:.1f}톤({K('S2_d27')/1000:.1f}톤) · S3 50%·안정화 3개월 → {K('S3_p')/1000:.1f}톤({K('S3_d27')/1000:.1f}톤). 개선 안정 후 47.2톤÷12 = 3,933 kg/월은 상한 참고 → 2027 대응에는 선행 생산·현재고 확인이 핵심.
■ 필요 Batch(출하÷190): 1~2월 14.5 / 3~6월 17.5 / 7~12월 21.5 Batch/월, 연 {K('nb27'):.1f} Batch (동일 Batch 다고객 배분 — 시간 중복 계산 안 함).
■ 충진: 5 Gal (하이닉스+CXMT)÷20×2 h = 189.3 / 247.3 / 267.3 h/월, 연 2,972 h (OQC 별도 연 {K('ac27'):.0f} h). 월평균 병 수는 병 상당 — 실제는 20 kg 정수 용기 배분. 200 L: 이지켐 4→8용기, 한솔 2용기 → 6 / 10용기/월, 연 96용기(이지켐 72 + 한솔 24), OQC 연 {K('al27'):.0f} h. 한솔 1~6월 수동·7월~ ARS, 이지켐 ARS 전환 월 확인 — 충진시간 미확인 고객은 확정하지 않음(공통 8 h 적용 가정 시 연 {K('ak27'):.0f} h, 비교용).
■ 재고: 기말재고 = 기초재고 + 출하 가능 양품 생산량 − 출하량 / 필요 생산량 = 출하량 + 목표 기말재고 − 기초재고. 하이닉스 출하·목표재고 우선 → 나머지 고객 배분(우선순위 미확인 시 출하 비율). 2027.12 기말재고 → 2028.1 기초재고 연결.""")

d = year_slide(s5, 2028)
set_lines(d["TextBox 2"], ["2028년 월별 출하 가정과 Capa. 대응"])
set_lines(d["TextBox 5"], ["출하 가정 49.12톤 (2027 하반기 수준 유지) · 47.2톤 대비 1.92톤 · 105℃ 50.2톤 검토"])
set_lines(d["TextBox 165"], [["① 2028 월별 출하 가정 · 참고 Capa. · 충진 (kg)", "   2027년 하반기 수준 유지 가정 — 확정 고객 수요 아님"]])
set_lines(d["TextBox 167"], ["* 하이닉스 22,720÷12 · CXMT 780 · 이지켐 1,120 · 한솔 300 kg/월 (2027 하반기 수준 유지 가정 — 2027 연간 43.4톤 복제 아님) · 105℃는 미확정 검토안으로 별도 비교 · 실제 생산량은 연간 환산값과 구분"])
set_lines(d["TextBox 169"], [["② 2028 연간 환산 비교", "   확정 부족량·판매 가능량 아님"]] if len(d["TextBox 169"].text_frame.paragraphs[0].runs) > 1 else ["② 2028 연간 환산 비교"])
kp5 = [("출하 가정 (하반기 수준 유지)", ["49.12", " 톤"], "4,093.3 kg/월 × 12 · 확정 수요 아님"),
       ("기본: 47.2톤 기준", [f"{K('d28_base'):.2f}".replace("-", "−"), " 톤"], "부족 방향 · 확정 부족량 아님"),
       ("105℃ 12개월 적용 시 (50.2)", [f"+{K('d28_105'):.2f}", " 톤"], "여유 방향 · 검토안 · 승인·시점 미정"),
       ("24개월 필요 선행재고 (참고)", [f"{K('S2_inv')/1000:.1f} ~ {K('S3_inv')/1000:.1f}", " 톤"], "'27.1.1 기준 · 시나리오 S2~S3 (Excel 05)")]
for (a, b_, c_), base in zip(kp5, (171, 175, 179, 183)):
    set_lines(d[f"TextBox {base}"], [a]); set_lines(d[f"TextBox {base+1}"], [b_]); set_lines(d[f"TextBox {base+2}"], [c_])
for base in (176, 180):
    for r in d[f"TextBox {base}"].text_frame.paragraphs[0].runs:
        r.font.color.rgb = RGBColor.from_string("C00000" if base == 176 else "385723")
set_lines(d["TextBox 187"], ["③ 재고 이월 · 추가 개선 · 품질 · 충진 자원 검토"])
kp5b = [("재고 이월", ["'27.12 → '28.1", ""], "기초·목표재고 입력 · 하이닉스 우선 확보"),
        ("추가 개선", [f"설비별 ≤ {K('rf_int_each'):.1f}", " h"], "Mix·준비·대기 단축 · 105℃ 검토"),
        ("충진 자원", ["267.3 h · 10", " 용기/월"], f"5 Gal 연 {K('gb28'):,.0f} h (1교대 {K('gb_ld1')*100:.0f}%) · 200 L 연 {K('c200_28'):.0f}용기"),
        ("105℃ 검토안 확인", ["승인 · 적용 시점", ""], "열 안정성·Dimer·Unknown·Yield · 이지켐 색도")]
for (a, b_, c_), base in zip(kp5b, (189, 193, 197, 201)):
    set_lines(d[f"TextBox {base}"], [a]); set_lines(d[f"TextBox {base+1}"], [[t for t in b_ if t]]); set_lines(d[f"TextBox {base+2}"], [c_])
notes(s5, f"""[5장 2028년 월별 출하 가정과 Capa. 대응 (Excel 01_Inputs K, 04_Monthly_2026_2028, 05_Reflux_Scenarios 3·5절)]
■ 출하 가정: 2027년 하반기 월 출하량을 12개월 유지하는 가정(확정 고객 수요 아님) — 하이닉스 22,720÷12, CXMT 780, 이지켐 1,120, 한솔 300 kg/월 → 4,093.33 kg/월 × 12 = 49,120 kg = 49.12톤 (CXMT 9,360 · 이지켐 13,440 · 한솔 3,600). 2027 연간 43.4톤을 그대로 복제하지 않음. 01_Inputs K에서 독립 변경 가능.
■ 연간 환산 비교(확정 부족량·판매 가능량 아님): 기본 47.2톤 기준 47.2 − 49.12 = −1.92톤(부족 방향) · 105℃ 검토안이 12개월 적용되는 경우 50.2 − 49.12 = +1.08톤(여유 방향). 105℃는 미확정·적용 시점 미정·검증 및 승인 필요 → 기본 시나리오는 47.2톤, 적용 시점 확정 시 01_Inputs '105℃ 적용 시작 월' 입력 → 04 시트 월별 반영.
■ 필요 Batch: 21.5 Batch/월, 연 {K('nb28'):.1f} Batch(190 kg). 47.2톤 환산 20.7 Batch/월. 설비별 등가 간격 {K('rf_int_each'):.1f} h 이하·Batch량·양품률이 대응 조건.
■ 충진: 5 Gal 267.3 h/월, 연 3,208 h (OQC 별도 연 {K('ac28'):.0f} h) · 200 L 이지켐 8 + 한솔 2 = 10용기/월, 연 120용기(96 + 24), 전월 ARS — ARS 충진시간 미확인(공통 8 h 적용 가정 시 연 {K('ak28'):.0f} h). 정제기 720 h를 인력 근무시간으로 적용하지 않음.
■ 재고: 2027.12 기말재고 → 2028.1 기초재고 연결(미확인 시 0 초기화 안 함). 시나리오 기준 2027.1.1 필요 선행재고 S1 {K('S1_inv')/1000:.1f} / S2 {K('S2_inv')/1000:.1f} / S3 {K('S3_inv')/1000:.1f}톤(24개월, 누적 최저 '28.12). 하이닉스 공급·목표재고 우선, 다른 고객 우선순위는 확인값.
■ 확인: 105℃ 열 안정성·Dimer·Unknown impurity·Yield 영향·고객 승인·적용 시점 / 이지켐 색도 합격률 / 글로브 박스 교대 / 2028 수요 확정.""")

prs.save(OUT)
print("saved", OUT)
